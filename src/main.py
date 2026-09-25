"""
main.py — Quintennial Data Pipeline (portfolio version)

Versão pública/sanitizada de uma ferramenta interna. A consulta ao banco
corporativo original (Oracle) foi substituída por um banco SQLite local com
dados 100% sintéticos (ver data/generate_synthetic_data.py e README.md).
As regras de negócio e o algoritmo de cálculo foram preservados.

Calcula quem completa quinquênios (5, 10, 15, 20... anos) em um mês/ano
informado pelo usuário.

REGRAS PRINCIPAIS
-----------------
- Permite agrupar as matrículas pelo CPF ou pelo PIS.
- Matrículas simultâneas não contam tempo em dobro.
- Hiatos sem vínculo não entram no tempo de serviço.
- A rescisão é uma data EXCLUSIVA: o dia da rescisão não é contado como
  dia trabalhado.
- Matrícula iniciada na mesma data da rescisão anterior é considerada
  continuidade, sem dia de hiato.
- O cálculo do quinquênio considera SOMENTE ANOS E MESES de serviço.
- A parcela de DIAS é deliberadamente ignorada na apuração do quinquênio.
- Para a contagem do quinquênio, cada intervalo contribui apenas com os meses
  completos no calendário (anos x 12 + meses), sem acumular dias restantes.
- A DATA-BASE é o último dia real do mês informado pelo usuário.
- A matrícula principal segue a regra já existente no código.
- A matrícula é tratada como texto e normalizada para 8 dígitos.
- CPF e PIS são tratados como texto, conforme o modo selecionado.

REGRAS DO CPF
-------------
- CPF com 11 dígitos:
      mantém exatamente os 11 dígitos.
- CPF vazio ou com quantidade diferente de 11 dígitos:
      rejeitado e enviado para a aba "Inválidos".
- Não há corte, preenchimento ou descarte automático de dígitos.
- O CPF normalizado é a CHAVE DE AGRUPAMENTO DA PESSOA.

VÍNCULO
-------
O vínculo NÃO é determinado pelo CPF.

A matrícula determina o prefixo:

      01 = ENTITY A
      02 = ENTITY B
      04 = ENTITY C

Porém, o campo de vínculo retornado pela consulta tem prioridade.
O prefixo da matrícula é usado como fallback caso a consulta não retorne
um vínculo.

IMPORTANTE — CPF x MATRÍCULA
----------------------------
Exemplo (valores fictícios):

CPF 12345678901

    Matrícula 01000001 -> prefixo 01 -> ENTITY A
    Matrícula 02000001 -> prefixo 02 -> ENTITY B
    Matrícula 04000001 -> prefixo 04 -> ENTITY C

Todas as matrículas podem pertencer ao mesmo histórico de pessoa quando
possuem o mesmo CPF normalizado.

O tempo é consolidado sobre os intervalos dessas matrículas, sem contagem
dupla em períodos sobrepostos.

ABA "INVÁLIDOS"
---------------
Todos os registros rejeitados durante a preparação dos dados são enviados
para uma aba separada chamada "Inválidos", permitindo identificar o motivo
exato e corrigir o cadastro na origem.

REGRAS DO PIS
-------------
- PIS com 11 dígitos: mantém exatamente os 11 dígitos.
- PIS com 12 dígitos: descarta sempre o primeiro dígito.
- PIS vazio ou com quantidade diferente de 11/12 dígitos: rejeitado e enviado
  para a aba "Inválidos".
- O PIS normalizado é a chave de agrupamento quando o modo PIS é selecionado.

REQUISITOS
----------
pip install -r requirements.txt
"""

import os
import re
import calendar
import sqlite3
import threading

from datetime import datetime, date, timedelta

import pandas as pd

from dateutil.relativedelta import relativedelta

from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment,
)

from openpyxl.utils import get_column_letter

from openpyxl.worksheet.page import PageMargins

import tkinter as tk

from tkinter import (
    ttk,
    filedialog,
    messagebox,
)


# ===========================================================================
# CONFIGURAÇÃO GLOBAL
# ===========================================================================

BASE_DIR = os.path.dirname(
    os.path.dirname(
        os.path.abspath(__file__)
    )
)


# ---------------------------------------------------------------------------
# FONTE DE DADOS
# ---------------------------------------------------------------------------
#
# Na versão de produção, os dados vinham de um banco Oracle corporativo
# (a conexão corporativa foi removida desta versão pública).
# Aqui a fonte é um banco SQLite local gerado por
# data/generate_synthetic_data.py. O caminho pode ser trocado pela variável
# de ambiente QUINTENNIAL_DB. Nenhuma credencial é necessária.

DB_PATH = os.environ.get(
    "QUINTENNIAL_DB",
    os.path.join(
        BASE_DIR,
        "data",
        "synthetic_hr.db",
    ),
)


# ---------------------------------------------------------------------------
# SAÍDA
# ---------------------------------------------------------------------------

PASTA_SAIDA_PADRAO = os.path.join(
    BASE_DIR,
    "output",
)


# ---------------------------------------------------------------------------
# REGRAS DE NEGÓCIO
# ---------------------------------------------------------------------------

ANOS_POR_QUINQUENIO = 5
MAX_MULTIPLOS = 30
LARGURA_MATRICULA = 8


VINCULOS_POR_PREFIXO = {
    "01": "ENTITY A",
    "02": "ENTITY B",
    "04": "ENTITY C",
}


# ---------------------------------------------------------------------------
# CORES
# ---------------------------------------------------------------------------

azul_escuro = "1F4E78"
azul = "5B9BD5"
azul_claro = "D9EAF7"
azul_muito_claro = "EEF5FB"
verde = "70AD47"
verde_claro = "E2F0D9"
amarelo_claro = "FFF2CC"
cinza_grade = "B7C9D6"
branco = "FFFFFF"

cinza = "666666"
vermelho = "C00000"
vermelho_claro = "FCE4D6"

BG_MAIN = "#EAF4FC"
BG_CARD = "#FFFFFF"
BG_SOFT = "#DCEBFA"
BG_MENU = "#2F5D8C"

ACCENT = "#6FA8DC"
ACCENT_DARK = "#3D6FA8"
ACCENT_HOVER = "#5A93C7"
ACCENT_SUCCESS = "#2E7D32"
ACCENT_DANGER = "#C62828"

TEXT_MAIN = "#274B6D"
TEXT_MUTED = "#6C88A3"

FONT = "Segoe UI"


# ---------------------------------------------------------------------------
# COLUNAS ESPERADAS
# ---------------------------------------------------------------------------

COLUNAS_ESPERADAS = {
    "MATRÍCULA": "matricula",
    "NOME_FUNCIONÁRIO": "nome",
    "CPF": "cpf",
    "PIS": "pis",
    "FUNÇÃO": "funcao",
    "LOTAÇÃO": "lotacao",
    "JORNADA": "jornada",
    "ADMISSÃO": "admissao",
    "RESCISÃO": "rescisao",
    "EMPRESA": "empresa",
    "VÍNCULO": "vinculo_original",
}



# ===========================================================================
# FUNÇÕES AUXILIARES
# ===========================================================================

def limpar_texto_excel(valor):
    """
    Remove caracteres de controle que o Excel/OpenPyXL não aceita.
    """

    if isinstance(valor, str):
        return re.sub(
            r"[\x00-\x08\x0B\x0C\x0E-\x1F]",
            "",
            valor,
        )

    return valor


def formula_excel_como_texto(formula):
    """
    Guarda uma fórmula equivalente do Excel como TEXTO.

    O apostrofo inicial impede que o Excel tente executá-la.
    """

    formula = str(formula).strip()

    if not formula.startswith("="):
        formula = "=" + formula

    return "'" + formula


# ===========================================================================
# CPF
# ===========================================================================

def diagnosticar_cpf(valor):
    """
    Diagnóstico completo do CPF.

    Regras:
        CPF vazio -> rejeita.
        CPF com 11 dígitos -> mantém exatamente os 11 dígitos.
        Qualquer outra quantidade de dígitos -> rejeita.
    """

    if pd.isna(valor):
        return {
            "original": "",
            "digitos": "",
            "normalizado": None,
            "situacao": "CPF VAZIO",
        }

    original = str(valor).strip()
    digitos = re.sub(r"\D", "", original)

    if not digitos:
        return {
            "original": original,
            "digitos": "",
            "normalizado": None,
            "situacao": "SEM DÍGITOS",
        }

    if len(digitos) == 11:
        return {
            "original": original,
            "digitos": digitos,
            "normalizado": digitos,
            "situacao": "11 DÍGITOS — MANTIDO",
        }

    return {
        "original": original,
        "digitos": digitos,
        "normalizado": None,
        "situacao": f"{len(digitos)} DÍGITOS — REJEITADO",
    }


def limpar_cpf(valor):
    """Normaliza o CPF: aceita somente 11 dígitos."""
    diagnostico = diagnosticar_cpf(valor)
    return diagnostico["normalizado"]


# ===========================================================================
# PIS
# ===========================================================================

def diagnosticar_pis(valor):
    """
    Diagnóstico completo do PIS.

    Regras:
        PIS vazio -> rejeita.
        11 dígitos -> mantém.
        12 dígitos -> remove o primeiro.
        demais quantidades -> rejeita.
    """

    if pd.isna(valor):
        return {
            "original": "",
            "digitos": "",
            "normalizado": None,
            "situacao": "PIS VAZIO",
        }

    original = str(valor).strip()
    if original == "":
        return {
            "original": "",
            "digitos": "",
            "normalizado": None,
            "situacao": "PIS VAZIO",
        }

    digitos = re.sub(r"\D", "", original)

    if not digitos:
        return {
            "original": original,
            "digitos": "",
            "normalizado": None,
            "situacao": "SEM DÍGITOS",
        }

    if len(digitos) == 11:
        return {
            "original": original,
            "digitos": digitos,
            "normalizado": digitos,
            "situacao": "11 DÍGITOS — MANTIDO",
        }

    if len(digitos) == 12:
        return {
            "original": original,
            "digitos": digitos,
            "normalizado": digitos[1:],
            "situacao": "12 DÍGITOS — 1º DESCARTADO",
        }

    return {
        "original": original,
        "digitos": digitos,
        "normalizado": None,
        "situacao": f"{len(digitos)} DÍGITOS — REJEITADO",
    }


def limpar_pis(valor):
    """Normaliza o PIS usando as regras oficiais do programa."""
    diagnostico = diagnosticar_pis(valor)
    return diagnostico["normalizado"]



# ===========================================================================
# MATRÍCULA
# ===========================================================================

def normalizar_matricula(valor):
    """
    Mantém matrícula como texto e recupera zeros à esquerda.

    Exemplo:

        1000273 -> 01000273
    """

    if pd.isna(valor):
        return ""

    texto = str(valor).strip()

    if (
        texto.endswith(".0")
        and texto[:-2].isdigit()
    ):
        texto = texto[:-2]

    if texto.isdigit():
        texto = texto.zfill(
            LARGURA_MATRICULA
        )

    return texto


def prefixo_matricula(matricula):
    """
    Retorna os dois primeiros caracteres da matrícula.
    """

    matricula = normalizar_matricula(
        matricula
    )

    return matricula[:2]


def vinculo_por_matricula(matricula):
    """
    Determina vínculo pelo prefixo da matrícula.
    """

    prefixo = prefixo_matricula(
        matricula
    )

    return VINCULOS_POR_PREFIXO.get(
        prefixo,
        "NÃO IDENTIFICADO",
    )


# ===========================================================================
# DATAS
# ===========================================================================

def parse_data(valor):

    if valor is None:
        return None

    if (
        isinstance(valor, float)
        and pd.isna(valor)
    ):
        return None

    if isinstance(valor, datetime):
        return valor.date()

    if isinstance(valor, date):
        return valor

    if isinstance(valor, pd.Timestamp):
        return valor.date()

    texto = str(valor).strip()

    if (
        texto == ""
        or texto.upper()
        in (
            "ATIVO",
            "ATIVA",
            "NAN",
            "NAT",
        )
    ):
        return None

    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
    ):

        try:

            return datetime.strptime(
                texto,
                fmt,
            ).date()

        except ValueError:
            pass

    for fmt in (
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%d/%m/%y",
    ):

        try:

            return datetime.strptime(
                texto,
                fmt,
            ).date()

        except ValueError:
            pass

    try:

        return pd.to_datetime(
            texto,
            dayfirst=True,
        ).date()

    except Exception:

        return None


def eh_ativo(valor_rescisao):

    if valor_rescisao is None:
        return True

    if (
        isinstance(valor_rescisao, float)
        and pd.isna(valor_rescisao)
    ):
        return True

    return (
        isinstance(
            valor_rescisao,
            str,
        )
        and valor_rescisao.strip().upper()
        in (
            "ATIVO",
            "ATIVA",
        )
    )


def ultimo_dia_do_mes(
    mes,
    ano,
):

    return date(
        ano,
        mes,
        calendar.monthrange(
            ano,
            mes,
        )[1],
    )


# ===========================================================================
# TEMPO
# ===========================================================================

def duracao_calendario(
    inicio,
    fim,
):

    if fim <= inicio:
        return relativedelta()

    return relativedelta(
        fim,
        inicio,
    )


def meses_intervalo(
    inicio,
    fim,
):
    """
    Retorna a quantidade de meses de serviço considerando
    EXCLUSIVAMENTE ANO e MÊS.

    REGRA FUNDAMENTAL
    -----------------
    O DIA é ignorado.

    Porém, o MÊS e o ANO NÃO são ignorados.

    Portanto:

        15/06/2012 -> 14/06/2026
        = 168 meses
        = 14 anos

    mas:

        01/11/2021 -> 30/10/2026
        = 59 meses
        = 4 anos e 11 meses

    e somente:

        01/11/2021 -> 30/11/2026
        = 60 meses
        = 5 anos

    Ou seja:

        DIA      -> ignorado
        MÊS/ANO  -> obrigatoriamente respeitado

    A data-base também funciona como limite rígido.
    Nenhum período pode ser projetado além dela.
    """

    if (
        inicio is None
        or fim is None
        or fim <= inicio
    ):
        return 0

    total_meses = (
        (fim.year - inicio.year) * 12
        + (fim.month - inicio.month)
    )

    return max(
        0,
        total_meses,
    )
def formatar_meses(
    total_meses,
):
    """
    Formata tempo exclusivamente em anos e meses.
    """
    total_meses = max(
        0,
        int(total_meses),
    )

    anos = total_meses // 12
    meses = total_meses % 12

    return f"{anos} anos e {meses} meses"


def formatar_relativedelta(rd):

    return formatar_meses(
        rd.years * 12 + rd.months
    )


def mesclar_intervalos(intervalos):
    """
    Mescla períodos sobrepostos ou sem hiato real.

    Rescisão é exclusiva.
    """

    intervalos = sorted(
        intervalos,
        key=lambda i: i["inicio"],
    )

    mesclados = []

    for intervalo in intervalos:

        atual = dict(
            intervalo
        )

        if not mesclados:

            mesclados.append(
                atual
            )

            continue

        ultimo = mesclados[-1]

        if ultimo["aberto"]:

            continue

        if atual["inicio"] <= ultimo["fim"]:

            if atual["aberto"]:

                ultimo["aberto"] = True
                ultimo["fim"] = None

            else:

                ultimo["fim"] = max(
                    ultimo["fim"],
                    atual["fim"],
                )

        else:

            mesclados.append(
                atual
            )

    return mesclados


def calcular_meses_totais(
    intervalos,
    data_limite,
):
    """
    Soma o tempo de serviço em meses completos, ignorando dias.
    """
    total_meses = 0

    for intervalo in sorted(
        intervalos,
        key=lambda x: x["inicio"],
    ):

        inicio = intervalo["inicio"]

        if inicio >= data_limite:
            continue

        fim = (
            data_limite
            if intervalo["aberto"]
            else intervalo["fim"]
        )

        if fim is None:
            fim = data_limite

        if fim > data_limite:
            fim = data_limite

        if fim <= inicio:
            continue

        total_meses += meses_intervalo(
            inicio,
            fim,
        )

    return total_meses


def formatar_tempo(
    intervalos,
    data_limite,
):
    """
    Retorna o tempo total somente em anos e meses.

    Dias restantes dos intervalos são descartados e NÃO são carregados
    para formar meses posteriores.
    """
    return formatar_meses(
        calcular_meses_totais(
            intervalos,
            data_limite,
        )
    )


def anos_completos(
    intervalos,
    data_limite,
):
    """
    Retorna somente os anos completos a partir dos meses completos.
    """
    total_meses = calcular_meses_totais(
        intervalos,
        data_limite,
    )

    return total_meses // 12


def meses_completos(
    intervalos,
    data_limite,
):
    """
    Retorna o total de meses completos de serviço.
    """
    return calcular_meses_totais(
        intervalos,
        data_limite,
    )


def virtual_fim_apos_intervalo(
    cursor_virtual,
    inicio,
    fim,
):

    if fim <= inicio:
        return cursor_virtual

    return (
        cursor_virtual
        + duracao_calendario(
            inicio,
            fim,
        )
    )


def calcular_datas_quinquenio(
    intervalos_mesclados,
    data_base,
):
    """
    Calcula os quinquênios usando SOMENTE ANOS e MESES.

    Cada intervalo contribui com:

        (ano_final - ano_inicial) * 12
        +
        (mes_final - mes_inicial)

    O DIA é ignorado.

    ATENÇÃO:
    ignorar o dia NÃO significa ignorar o mês.

    Exemplo:

        01/11/2021 -> 30/10/2026
        = 59 meses
        = não alcança 60 meses

        01/11/2021 -> 30/11/2026
        = 60 meses
        = alcança 5 anos

    A data-base é um limite rígido.
    Nenhum quinquênio pode ser criado depois
    do mês/ano efetivamente disponível no intervalo.
    """

    resultados = []

    intervalos = sorted(
        intervalos_mesclados,
        key=lambda i: i["inicio"],
    )

    meses_acumulados = 0
    proximo_multiplo = 1

    for intervalo in intervalos:

        inicio = intervalo["inicio"]

        # --------------------------------------------------------------
        # FIM EFETIVO
        # --------------------------------------------------------------

        fim = (
            data_base
            if intervalo["aberto"]
            else intervalo["fim"]
        )

        if fim is None:
            fim = data_base

        # Nunca ultrapassar a data-base.
        if fim > data_base:
            fim = data_base

        if (
            inicio >= data_base
            or fim <= inicio
        ):
            continue

        # --------------------------------------------------------------
        # MESES DISPONÍVEIS
        # --------------------------------------------------------------

        meses_disponiveis = meses_intervalo(
            inicio,
            fim,
        )

        if meses_disponiveis <= 0:
            continue

        meses_antes = meses_acumulados

        meses_acumulados += (
            meses_disponiveis
        )

        # --------------------------------------------------------------
        # MARCOS DE 60, 120, 180...
        # --------------------------------------------------------------

        while proximo_multiplo <= MAX_MULTIPLOS:

            meta_meses = (
                proximo_multiplo
                * ANOS_POR_QUINQUENIO
                * 12
            )

            # Ainda não atingiu o próximo marco.
            if meta_meses > meses_acumulados:
                break

            meses_necessarios_no_intervalo = (
                meta_meses
                - meses_antes
            )

            if (
                meses_necessarios_no_intervalo
                <= 0
            ):
                proximo_multiplo += 1
                continue

            # ----------------------------------------------------------
            # DATA DO MARCO
            # ----------------------------------------------------------
            #
            # O relativedelta é utilizado somente para localizar
            # o mês/ano correspondente ao marco.
            #
            # O dia não participa da decisão final.
            # ----------------------------------------------------------

            data_atingida = (
                inicio
                + relativedelta(
                    months=(
                        meses_necessarios_no_intervalo
                    )
                )
            )

            # ----------------------------------------------------------
            # BLINDAGEM DA DATA-BASE
            # ----------------------------------------------------------
            #
            # O marco precisa estar dentro do período efetivamente
            # disponível, considerando ANO + MÊS.
            #
            # O dia continua ignorado.
            # ----------------------------------------------------------

            marco_mes_ano = (
                data_atingida.year,
                data_atingida.month,
            )

            limite_mes_ano = (
                fim.year,
                fim.month,
            )

            if marco_mes_ano > limite_mes_ano:
                break

            # ----------------------------------------------------------
            # BLINDAGEM EXTRA DA DATA-BASE
            # ----------------------------------------------------------
            #
            # Mesmo com a comparação acima, não permitimos que um
            # resultado seja criado fora do mês/ano da data-base.
            # ----------------------------------------------------------

            data_base_mes_ano = (
                data_base.year,
                data_base.month,
            )

            if marco_mes_ano > data_base_mes_ano:
                break

            resultados.append(
                (
                    proximo_multiplo
                    * ANOS_POR_QUINQUENIO,
                    data_atingida,
                )
            )

            proximo_multiplo += 1

        if (
            proximo_multiplo
            > MAX_MULTIPLOS
        ):
            break

    return resultados

def selecionar_matricula_principal(
    registros,
):

    ativos = [
        r
        for r in registros
        if r["aberto"]
    ]

    candidatos = (
        ativos
        if ativos
        else registros
    )

    return max(
        candidatos,
        key=lambda r: r["inicio"],
    )


def dias_intervalo(
    inicio,
    fim,
):

    if fim <= inicio:
        return 0

    return (
        fim - inicio
    ).days


# ===========================================================================
# SQL
# ===========================================================================

SQL_QUINQUENIOS = """
SELECT EMP.badge_id AS "MATRÍCULA",
       EMP.full_name AS "NOME_FUNCIONÁRIO",
       ATTR_TAX.value AS "CPF",
       EMP.pis AS "PIS",
       ATTR_JOB.value AS "FUNÇÃO",
       DEP.description AS "LOTAÇÃO",
       EMP.daily_minutes / 60.0 * 5 AS "JORNADA",
       EMP.hire_date AS "ADMISSÃO",
       CASE
           WHEN EMP.termination_date = '1899-12-30' THEN 'ATIVO'
           ELSE strftime('%d/%m/%Y', EMP.termination_date)
       END AS "RESCISÃO",
       EMP.company_code AS "EMPRESA",
       CASE EMP.employer_code
           WHEN 'ENT-A'  THEN 'ENTITY A'
           WHEN 'ENT-B'  THEN 'ENTITY B'
           WHEN 'ENT-C1' THEN 'ENTITY C (CLT)'
           WHEN 'ENT-C2' THEN 'ENTITY C (COMMISSIONED)'
           WHEN 'ENT-X'  THEN 'APPRENTICE/OTHER'
           ELSE 'OTHER'
       END AS "VÍNCULO"
FROM employee EMP
INNER JOIN department DEP
        ON (EMP.department_id = DEP.department_id
            AND EMP.company_code = DEP.company_code)
INNER JOIN employee_attribute ATTR_TAX
        ON (ATTR_TAX.badge_id = EMP.badge_id
            AND ATTR_TAX.company_code = EMP.company_code
            AND ATTR_TAX.attribute_code = 'TAX_ID')
INNER JOIN employee_attribute ATTR_JOB
        ON (ATTR_JOB.badge_id = EMP.badge_id
            AND ATTR_JOB.company_code = EMP.company_code
            AND ATTR_JOB.attribute_code = 'JOB_TITLE')
WHERE EMP.company_code IN ('CO1', 'CO2')
  AND EMP.badge_id NOT IN ('99999901', '99999902')
ORDER BY EMP.full_name
"""

# Notas sobre a consulta original (produção, dialeto Oracle):
#   - a data de rescisão "em aberto" era representada por uma data sentinela
#     (30/12/1899), convertida em 'ATIVO' via DECODE; aqui isso é um CASE;
#   - CPF e função vinham de uma tabela de atributos por código (duas junções
#     na mesma tabela), preservado aqui em employee_attribute;
#   - um pequeno conjunto de matrículas técnicas era excluído por NOT IN;
#   - nomes de tabelas, colunas e códigos internos foram substituídos por um
#     esquema fictício (ver data/generate_synthetic_data.py).


# ===========================================================================
# BANCO DE DADOS (SQLite sintético)
# ===========================================================================

def conectar_banco():
    """
    Abre o banco SQLite sintético em modo somente leitura.

    Substitui a conexão Oracle da versão de produção.
    """

    caminho = os.path.abspath(
        DB_PATH
    )

    if not os.path.isfile(
        caminho
    ):

        raise FileNotFoundError(
            "Banco sintético não encontrado:\n"
            f"{caminho}\n\n"
            "Gere-o com:\n"
            "python data/generate_synthetic_data.py"
        )

    try:

        return sqlite3.connect(
            f"file:{caminho.replace(os.sep, '/')}?mode=ro",
            uri=True,
        )

    except Exception as exc:

        raise RuntimeError(
            f"Falha ao abrir o banco de dados:\n{exc}"
        ) from exc


def consultar_banco(
    logger=None,
):

    conn = conectar_banco()

    try:

        if logger:

            logger(
                "Executando consulta no banco de dados..."
            )

        cursor = conn.cursor()

        cursor.execute(
            SQL_QUINQUENIOS
        )

        colunas = [
            desc[0]
            for desc in cursor.description
        ]

        linhas = cursor.fetchall()

        df = pd.DataFrame(
            linhas,
            columns=colunas,
        )

        if logger:

            logger(
                "Consulta concluída: "
                f"{len(df):,} registros retornados."
            )

        return df

    finally:

        conn.close()


# ===========================================================================
# DATAFRAME
# ===========================================================================

def validar_dataframe(df):

    colunas = [
        str(c).strip()
        for c in df.columns
    ]

    df = df.copy()

    df.columns = colunas

    faltando = [
        c
        for c in COLUNAS_ESPERADAS
        if c not in df.columns
    ]

    if faltando:

        raise ValueError(
            "Colunas esperadas não encontradas "
            "no resultado da consulta:\n"
            + "\n".join(
                f"- {c}"
                for c in faltando
            )
        )

    df = df.rename(
        columns=COLUNAS_ESPERADAS
    )

    try:

        df = df.map(
            limpar_texto_excel
        )

    except AttributeError:

        df = df.apply(
            lambda coluna:
            coluna.map(
                limpar_texto_excel
            )
        )

    return df


def preparar_registros(
    df,
    data_base,
    logger=None,
    modo="CPF",
):

    modo = str(modo).strip().upper()

    if modo not in ("CPF", "PIS"):
        raise ValueError(
            "Modo de agrupamento inválido. Use CPF ou PIS."
        )

    df = validar_dataframe(df)

    df["matricula"] = (
        df["matricula"].apply(
            normalizar_matricula
        )
    )

    linhas_invalidas = 0
    registros_invalidos = []
    registros_por_cpf = {}

    chave_campo = "cpf" if modo == "CPF" else "pis"
    diagnosticar_chave = (
        diagnosticar_cpf
        if modo == "CPF"
        else diagnosticar_pis
    )

    for _, linha in df.iterrows():

        # ==============================================================
        # CHAVE DE AGRUPAMENTO
        # ==============================================================

        chave_original = linha[chave_campo]

        diagnostico = diagnosticar_chave(
            chave_original
        )

        chave = diagnostico[
            "normalizado"
        ]

        # Mantemos os nomes internos "cpf_*" por compatibilidade com
        # as abas existentes; no modo PIS eles representam a chave PIS.
        cpf_original = chave_original
        cpf = chave

        # ==============================================================
        # MATRÍCULA
        # ==============================================================

        matricula_original = linha["matricula"]

        matricula = normalizar_matricula(
            matricula_original
        )

        # ==============================================================
        # ADMISSÃO
        # ==============================================================

        admissao = parse_data(
            linha["admissao"]
        )

        # ==============================================================
        # RESCISÃO
        # ==============================================================

        ativo = eh_ativo(
            linha["rescisao"]
        )

        rescisao = (
            None
            if ativo
            else parse_data(
                linha["rescisao"]
            )
        )

        # ==============================================================
        # MOTIVOS
        # ==============================================================

        motivos = []

        if cpf is None:
            motivos.append(
                diagnostico["situacao"]
            )

        if not matricula:
            motivos.append(
                "MATRÍCULA VAZIA / INVÁLIDA"
            )

        if admissao is None:
            motivos.append(
                "ADMISSÃO VAZIA / INVÁLIDA"
            )

        if (
            not ativo
            and rescisao is None
        ):
            valor_rescisao = str(
                linha["rescisao"]
            ).strip()

            motivos.append(
                "RESCISÃO VAZIA / INVÁLIDA "
                f"({valor_rescisao or 'vazio'})"
            )

        # ==============================================================
        # INVALIDADO
        # ==============================================================

        if motivos:

            linhas_invalidas += 1

            registros_invalidos.append(
                {
                    "nome": limpar_texto_excel(
                        linha.get(
                            "nome",
                            "",
                        )
                    ),
                    "cpf_original": limpar_texto_excel(
                        (
                            ""
                            if pd.isna(cpf_original)
                            else str(cpf_original)
                        )
                    ),
                    "cpf_digitos": diagnostico[
                        "digitos"
                    ],
                    "cpf_normalizado": cpf or "",
                    "cpf_situacao": diagnostico[
                        "situacao"
                    ],
                    "matricula_original": limpar_texto_excel(
                        (
                            ""
                            if pd.isna(matricula_original)
                            else str(matricula_original)
                        )
                    ),
                    "matricula_normalizada": matricula,
                    "prefixo": (
                        prefixo_matricula(matricula)
                        if matricula
                        else ""
                    ),
                    "admissao_original": limpar_texto_excel(
                        (
                            ""
                            if pd.isna(linha["admissao"])
                            else str(linha["admissao"])
                        )
                    ),
                    "admissao": (
                        admissao.strftime("%d/%m/%Y")
                        if admissao
                        else ""
                    ),
                    "rescisao_original": limpar_texto_excel(
                        (
                            ""
                            if pd.isna(linha["rescisao"])
                            else str(linha["rescisao"])
                        )
                    ),
                    "rescisao": (
                        "ATIVO"
                        if ativo
                        else (
                            rescisao.strftime("%d/%m/%Y")
                            if rescisao
                            else ""
                        )
                    ),
                    "motivos": " | ".join(motivos),
                    "empresa": limpar_texto_excel(
                        str(
                            linha.get(
                                "empresa",
                                "",
                            )
                            or ""
                        ).strip()
                    ),
                    "vinculo": limpar_texto_excel(
                        str(
                            linha.get(
                                "vinculo_original",
                                "",
                            )
                            or ""
                        ).strip()
                    ),
                    "funcao": limpar_texto_excel(
                        str(
                            linha.get(
                                "funcao",
                                "",
                            )
                            or ""
                        ).strip()
                    ),
                    "lotacao": limpar_texto_excel(
                        str(
                            linha.get(
                                "lotacao",
                                "",
                            )
                            or ""
                        ).strip()
                    ),
                    "jornada": limpar_texto_excel(
                        str(
                            linha.get(
                                "jornada",
                                "",
                            )
                            or ""
                        ).strip()
                    ),
                }
            )

            continue

        # ==============================================================
        # VÍNCULO
        # ==============================================================

        vinculo_fonte = str(
            linha.get(
                "vinculo_original",
                "",
            )
            or ""
        ).strip()

        vinculo = (
            vinculo_fonte
            or vinculo_por_matricula(
                matricula
            )
        )

        # ==============================================================
        # REGISTRO VÁLIDO
        # ==============================================================

        registro = {
            "nome": str(
                linha["nome"]
            ).strip(),
            "cpf": cpf,
            "cpf_original": (
                str(cpf_original).strip()
                if not pd.isna(cpf_original)
                else ""
            ),
            "cpf_situacao": diagnostico[
                "situacao"
            ],
            "cpf_real": (
                str(linha["cpf"]).strip()
                if not pd.isna(linha["cpf"])
                else ""
            ),
            "pis_real": (
                str(linha["pis"]).strip()
                if not pd.isna(linha["pis"])
                else ""
            ),
            "matricula": matricula,
            "prefixo": prefixo_matricula(
                matricula
            ),
            "inicio": admissao,
            "fim": rescisao,
            "aberto": ativo,
            "vinculo": vinculo,
            "vinculo_fonte": vinculo_fonte,
            "funcao": str(
                linha.get(
                    "funcao",
                    "",
                )
                or ""
            ).strip(),
            "lotacao": str(
                linha.get(
                    "lotacao",
                    "",
                )
                or ""
            ).strip(),
            "jornada": str(
                linha.get(
                    "jornada",
                    "",
                )
                or ""
            ).strip(),
            "empresa": str(
                linha.get(
                    "empresa",
                    "",
                )
                or ""
            ).strip(),
        }

        registros_por_cpf.setdefault(
            cpf,
            []
        ).append(registro)

    # ==============================================================
    # LOG
    # ==============================================================

    if logger:

        logger(
            f"{modo} válidos agrupados: "
            f"{len(registros_por_cpf):,} pessoas."
        )

        if linhas_invalidas:

            logger(
                "Aviso: "
                f"{linhas_invalidas:,} linha(s) "
                "ignoradas por dados inválidos."
            )

            contagem_motivos = {}

            for item in registros_invalidos:

                for motivo in item[
                    "motivos"
                ].split(" | "):

                    contagem_motivos[
                        motivo
                    ] = (
                        contagem_motivos.get(
                            motivo,
                            0
                        ) + 1
                    )

            logger(
                "Motivos encontrados:"
            )

            for motivo, quantidade in sorted(
                contagem_motivos.items(),
                key=lambda x: (
                    -x[1],
                    x[0],
                ),
            ):

                logger(
                    f"  - {motivo}: "
                    f"{quantidade:,}"
                )

        else:

            logger(
                "Nenhum registro inválido encontrado."
            )

    return (
        registros_por_cpf,
        registros_invalidos,
    )


# ===========================================================================
# ABA INVÁLIDOS
# ===========================================================================

def montar_aba_invalidos(
    workbook,
    registros_invalidos,
):

    borda = Border(
        left=Side(
            style="thin",
            color=cinza_grade,
        ),
        right=Side(
            style="thin",
            color=cinza_grade,
        ),
        top=Side(
            style="thin",
            color=cinza_grade,
        ),
        bottom=Side(
            style="thin",
            color=cinza_grade,
        ),
    )

    ws = workbook.create_sheet(
        "Inválidos"
    )

    ws.sheet_view.showGridLines = False

    # ==================================================================
    # TÍTULO
    # ==================================================================

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=14,
    )

    c = ws.cell(
        1,
        1,
        "REGISTROS INVÁLIDOS — "
        "AUDITORIA DO CADASTRO",
    )

    c.font = Font(
        size=15,
        bold=True,
        color=branco,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=vermelho,
    )

    c.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    ws.row_dimensions[
        1
    ].height = 30

    # ==================================================================
    # EXPLICAÇÃO
    # ==================================================================

    ws.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=14,
    )

    c = ws.cell(
        2,
        1,
        (
            "Estas linhas foram excluídas do cálculo porque "
            "apresentaram algum dado essencial inválido. "
            "Use esta aba para localizar e corrigir o cadastro na origem."
        ),
    )

    c.font = Font(
        size=10,
        color=vermelho,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=vermelho_claro,
    )

    c.alignment = Alignment(
        vertical="center",
        wrap_text=True,
    )

    c.border = borda

    # ==================================================================
    # CABEÇALHO
    # ==================================================================

    cabecalhos = [
        "Nome",
        "CPF Original",
        "CPF — Dígitos",
        "CPF Normalizado",
        "Situação CPF",
        "Matrícula Original",
        "Matrícula Normalizada",
        "Prefixo",
        "Admissão Original",
        "Admissão",
        "Rescisão Original",
        "Rescisão",
        "Motivo(s) da Exclusão",
        "Lotação",
    ]

    linha = 4

    for col, titulo in enumerate(
        cabecalhos,
        start=1,
    ):

        c = ws.cell(
            linha,
            col,
            titulo,
        )

        c.font = Font(
            bold=True,
            color=branco,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=vermelho,
        )

        c.border = borda

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    ws.row_dimensions[
        linha
    ].height = 32

    linha += 1

    # ==================================================================
    # DADOS
    # ==================================================================

    if not registros_invalidos:

        ws.merge_cells(
            start_row=linha,
            start_column=1,
            end_row=linha,
            end_column=14,
        )

        c = ws.cell(
            linha,
            1,
            "Nenhum registro inválido foi encontrado.",
        )

        c.font = Font(
            bold=True,
            color=verde,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=verde_claro,
        )

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        c.border = borda

    else:

        for idx, item in enumerate(
            registros_invalidos
        ):

            valores = [
                item["nome"],
                item["cpf_original"],
                item["cpf_digitos"],
                item["cpf_normalizado"],
                item["cpf_situacao"],
                item["matricula_original"],
                item["matricula_normalizada"],
                item["prefixo"],
                item["admissao_original"],
                item["admissao"],
                item["rescisao_original"],
                item["rescisao"],
                item["motivos"],
                item["lotacao"],
            ]

            for col, valor in enumerate(
                valores,
                start=1,
            ):

                c = ws.cell(
                    linha,
                    col,
                    valor,
                )

                c.border = borda

                c.alignment = Alignment(
                    horizontal=(
                        "left"
                        if col in (
                            1,
                            5,
                            13,
                            14,
                        )
                        else "center"
                    ),
                    vertical="center",
                    wrap_text=True,
                )

                c.fill = PatternFill(
                    "solid",
                    fgColor=(
                        branco
                        if idx % 2 == 0
                        else vermelho_claro
                    ),
                )

                if col in (
                    2,
                    3,
                    4,
                    6,
                    7,
                    8,
                ):

                    c.number_format = "@"

            ws.cell(
                linha,
                13,
            ).font = Font(
                bold=True,
                color=vermelho,
            )

            ws.cell(
                linha,
                13,
            ).fill = PatternFill(
                "solid",
                fgColor=vermelho_claro,
            )

            linha += 1

    # ==================================================================
    # FILTRO
    # ==================================================================

    if registros_invalidos:

        ws.auto_filter.ref = (
            f"A4:N{linha - 1}"
        )

        ws.freeze_panes = "A5"

    # ==================================================================
    # LARGURAS
    # ==================================================================

    larguras = {
        1: 38,
        2: 18,
        3: 16,
        4: 18,
        5: 28,
        6: 20,
        7: 20,
        8: 10,
        9: 20,
        10: 14,
        11: 22,
        12: 14,
        13: 48,
        14: 30,
    }

    for col, largura in (
        larguras.items()
    ):

        ws.column_dimensions[
            get_column_letter(col)
        ].width = largura

    # ==================================================================
    # IMPRESSÃO
    # ==================================================================

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.page_margins = PageMargins(
        left=0.3,
        right=0.3,
        top=0.5,
        bottom=0.5,
        header=0.2,
        footer=0.2,
    )

    ws.oddFooter.center.text = (
        "Página &P de &N"
    )

    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = "808080"


# ===========================================================================
# ABA DETALHE GERAL
# ===========================================================================

def montar_aba_detalhe_geral(
    workbook,
    mes_alvo,
    ano_alvo,
    data_base,
    data_geracao,
    detalhe_geral_por_pessoa,
):

    borda = Border(
        left=Side(
            style="thin",
            color=cinza_grade,
        ),
        right=Side(
            style="thin",
            color=cinza_grade,
        ),
        top=Side(
            style="thin",
            color=cinza_grade,
        ),
        bottom=Side(
            style="thin",
            color=cinza_grade,
        ),
    )

    ws = workbook.create_sheet(
        "Detalhe por Mat. Geral"
    )

    ws.sheet_view.showGridLines = False

    cabecalhos = [
        "Nome",
        "CPF",
        "Matrícula",
        "Prefixo",
        "Vínculo",
        "Lotação",
        "Jornada",
        "Admissão",
        "Rescisão",
        "Status",
        "Dias",
        "Tempo na Matrícula",
        "Mat. Principal?",
        "Tempo Total CPF",
        "Anos Completos CPF",
        f"Quinquênio em {mes_alvo:02d}/{ano_alvo}?",
    ]

    for col_idx, titulo in enumerate(
        cabecalhos,
        start=1,
    ):

        c = ws.cell(
            1,
            col_idx,
            titulo,
        )

        c.font = Font(
            name="Calibri",
            bold=True,
            size=10,
            color=branco,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=azul_escuro,
        )

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        c.border = borda

    ws.row_dimensions[
        1
    ].height = 34

    row_idx = 2

    for pessoa in detalhe_geral_por_pessoa:

        quinquenios_mes = (
            pessoa["quinquenios_mes"]
        )

        q_txt = (
            ", ".join(
                f"{m} anos — "
                f"{d.strftime('%d/%m/%Y')}"
                for m, d in quinquenios_mes
            )
            if quinquenios_mes
            else ""
        )

        for mat_idx, mat in enumerate(
            pessoa["matriculas"]
        ):

            bg = (
                azul_muito_claro
                if mat_idx % 2 == 0
                else branco
            )

            valores = [
                mat["nome"],
                pessoa["cpf"],
                mat["matricula"],
                mat["prefixo"],
                mat["vinculo"],
                mat["lotacao"],
                mat["jornada"],
                mat["admissao"],
                mat["rescisao"],
                (
                    "ATIVO"
                    if mat["rescisao"]
                    == "ATIVO"
                    else "RESCINDIDA"
                ),
                mat["dias"],
                mat["tempo"],
                (
                    "SIM"
                    if mat["principal"]
                    else ""
                ),
                pessoa["tempo_total"],
                pessoa["anos_total"],
                q_txt,
            ]

            for col_idx, valor in enumerate(
                valores,
                start=1,
            ):

                c = ws.cell(
                    row_idx,
                    col_idx,
                    valor,
                )

                c.border = borda

                c.alignment = Alignment(
                    horizontal=(
                        "left"
                        if col_idx in (
                            1,
                            5,
                            6,
                            12,
                            14,
                            16,
                        )
                        else "center"
                    ),
                    vertical="center",
                    wrap_text=True,
                )

                c.fill = PatternFill(
                    "solid",
                    fgColor=bg,
                )

                if col_idx in (
                    2,
                    3,
                ):

                    c.number_format = "@"

                if mat["principal"]:

                    c.font = Font(
                        bold=True,
                        color=azul_escuro,
                    )

                if (
                    col_idx == 10
                    and str(valor).upper()
                    == "ATIVO"
                ):

                    c.font = Font(
                        bold=True,
                        color=verde,
                    )

                    c.fill = PatternFill(
                        "solid",
                        fgColor=verde_claro,
                    )

                if (
                    col_idx == 13
                    and valor == "SIM"
                ):

                    c.font = Font(
                        bold=True,
                        color=azul_escuro,
                    )

                    c.fill = PatternFill(
                        "solid",
                        fgColor=azul_claro,
                    )

                if (
                    col_idx == 16
                    and valor
                ):

                    c.font = Font(
                        bold=True,
                        color="7F5000",
                    )

                    c.fill = PatternFill(
                        "solid",
                        fgColor=amarelo_claro,
                    )

            row_idx += 1

    larguras = {
        1: 40,
        2: 16,
        3: 14,
        4: 10,
        5: 18,
        6: 30,
        7: 10,
        8: 14,
        9: 14,
        10: 14,
        11: 10,
        12: 26,
        13: 14,
        14: 26,
        15: 18,
        16: 34,
    }

    for col, largura in (
        larguras.items()
    ):

        ws.column_dimensions[
            get_column_letter(col)
        ].width = largura

    ws.auto_filter.ref = (
        f"A1:P{row_idx - 1}"
    )

    ws.freeze_panes = "A2"

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.page_margins = PageMargins(
        left=0.3,
        right=0.3,
        top=0.5,
        bottom=0.5,
        header=0.2,
        footer=0.2,
    )

    ws.oddFooter.center.text = (
        "Página &P de &N"
    )

    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = "808080"


# ===========================================================================
# ABA DETALHE POR MATRÍCULA
# ===========================================================================

def montar_aba_detalhe_matricula(
    workbook,
    mes_alvo,
    ano_alvo,
    data_base,
    data_geracao,
    detalhe_por_pessoa,
):

    borda = Border(
        left=Side(
            style="thin",
            color=cinza_grade,
        ),
        right=Side(
            style="thin",
            color=cinza_grade,
        ),
        top=Side(
            style="thin",
            color=cinza_grade,
        ),
        bottom=Side(
            style="thin",
            color=cinza_grade,
        ),
    )

    ws = workbook.create_sheet(
        "Detalhe por Matricula"
    )

    ws.sheet_view.showGridLines = False

    linha_atual = 1

    # ==================================================================
    # TÍTULO
    # ==================================================================

    ws.merge_cells(
        start_row=linha_atual,
        start_column=1,
        end_row=linha_atual,
        end_column=9,
    )

    c = ws.cell(
        linha_atual,
        1,
        "RELATÓRIO DE QUINQUÊNIOS — "
        "DETALHAMENTO POR FUNCIONÁRIO",
    )

    c.font = Font(
        name="Calibri",
        size=16,
        bold=True,
        color=branco,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=azul_escuro,
    )

    c.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    ws.row_dimensions[
        linha_atual
    ].height = 32

    linha_atual += 1

    # ==================================================================
    # CABEÇALHO DO RELATÓRIO
    # ==================================================================

    ws.merge_cells(
        start_row=linha_atual,
        start_column=1,
        end_row=linha_atual,
        end_column=4,
    )

    c = ws.cell(
        linha_atual,
        1,
        f"DATA-BASE: "
        f"{data_base.strftime('%d/%m/%Y')}",
    )

    c.font = Font(
        size=11,
        bold=True,
        color="7F5000",
    )

    c.fill = PatternFill(
        "solid",
        fgColor=amarelo_claro,
    )

    c.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    c.border = borda

    ws.merge_cells(
        start_row=linha_atual,
        start_column=5,
        end_row=linha_atual,
        end_column=9,
    )

    c = ws.cell(
        linha_atual,
        5,
        f"MÊS/ANO: "
        f"{mes_alvo:02d}/{ano_alvo}"
        f"  |  GERADO EM: "
        f"{data_geracao.strftime('%d/%m/%Y')}",
    )

    c.font = Font(
        size=10,
        bold=True,
        color=azul_escuro,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=azul_claro,
    )

    c.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    c.border = borda

    ws.row_dimensions[
        linha_atual
    ].height = 25

    linha_atual += 2

    # ==================================================================
    # FUNCIONÁRIOS
    # ==================================================================

    for pessoa in detalhe_por_pessoa:

        principal = pessoa["principal"]

        # --------------------------------------------------------------
        # NOME
        # --------------------------------------------------------------

        ws.merge_cells(
            start_row=linha_atual,
            start_column=1,
            end_row=linha_atual,
            end_column=9,
        )

        c = ws.cell(
            linha_atual,
            1,
            f"FUNCIONÁRIO: {pessoa['nome']}",
        )

        c.font = Font(
            size=13,
            bold=True,
            color=branco,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=azul_escuro,
        )

        c.alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

        ws.row_dimensions[
            linha_atual
        ].height = 25

        linha_atual += 1

        # --------------------------------------------------------------
        # IDENTIFICAÇÃO 1
        # --------------------------------------------------------------

        blocos_linha1 = [
            (
                1,
                3,
                f"MATRÍCULA PRINCIPAL: "
                f"{principal['matricula']}",
            ),
            (
                4,
                6,
                f"CPF: {pessoa['cpf']}",
            ),
            (
                7,
                9,
                f"PREFIXO: "
                f"{principal['prefixo']}",
            ),
        ]

        for inicio, fim, texto in (
            blocos_linha1
        ):

            ws.merge_cells(
                start_row=linha_atual,
                start_column=inicio,
                end_row=linha_atual,
                end_column=fim,
            )

            c = ws.cell(
                linha_atual,
                inicio,
                texto,
            )

            c.font = Font(
                size=10,
                bold=True,
                color=azul_escuro,
            )

            c.fill = PatternFill(
                "solid",
                fgColor=azul_claro,
            )

            c.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            for col in range(
                inicio,
                fim + 1,
            ):

                ws.cell(
                    linha_atual,
                    col,
                ).border = borda

        linha_atual += 1

        # --------------------------------------------------------------
        # IDENTIFICAÇÃO 2
        # --------------------------------------------------------------

        blocos_linha2 = [
            (
                1,
                3,
                f"VÍNCULO: "
                f"{principal['vinculo']}",
            ),
            (
                4,
                6,
                f"STATUS: "
                f"{'ATIVO' if principal['aberto'] else 'INATIVO'}",
            ),
            (
                7,
                9,
                f"DATA-BASE: "
                f"{data_base.strftime('%d/%m/%Y')}",
            ),
        ]

        for inicio, fim, texto in (
            blocos_linha2
        ):

            ws.merge_cells(
                start_row=linha_atual,
                start_column=inicio,
                end_row=linha_atual,
                end_column=fim,
            )

            c = ws.cell(
                linha_atual,
                inicio,
                texto,
            )

            c.font = Font(
                size=10,
                bold=True,
                color=(
                    verde
                    if "ATIVO"
                    in texto
                    else azul_escuro
                ),
            )

            c.fill = PatternFill(
                "solid",
                fgColor=(
                    verde_claro
                    if "ATIVO"
                    in texto
                    else azul_claro
                ),
            )

            c.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            for col in range(
                inicio,
                fim + 1,
            ):

                ws.cell(
                    linha_atual,
                    col,
                ).border = borda

        linha_atual += 1

        # --------------------------------------------------------------
        # QUINQUÊNIO + TEMPO
        # --------------------------------------------------------------

        qtxt = ", ".join(
            f"{m} anos — "
            f"{d.strftime('%d/%m/%Y')}"
            for m, d in pessoa["quinquenios"]
        )

        ws.merge_cells(
            start_row=linha_atual,
            start_column=1,
            end_row=linha_atual,
            end_column=4,
        )

        c = ws.cell(
            linha_atual,
            1,
            f"QUINQUÊNIO GERADO: {qtxt}",
        )

        c.font = Font(
            size=11,
            bold=True,
            color="7F5000",
        )

        c.fill = PatternFill(
            "solid",
            fgColor=amarelo_claro,
        )

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        c.border = borda

        for col in range(
            1,
            5,
        ):

            ws.cell(
                linha_atual,
                col,
            ).border = borda

        ws.merge_cells(
            start_row=linha_atual,
            start_column=5,
            end_row=linha_atual,
            end_column=9,
        )

        c = ws.cell(
            linha_atual,
            5,
            f"TEMPO ATÉ A DATA-BASE: "
            f"{pessoa['tempo_total']} "
            f"({pessoa['anos_total']} "
            f"anos completos)",
        )

        c.font = Font(
            size=10,
            bold=True,
            color=verde,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=verde_claro,
        )

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        c.border = borda

        for col in range(
            5,
            10,
        ):

            ws.cell(
                linha_atual,
                col,
            ).border = borda

        linha_atual += 2

        # --------------------------------------------------------------
        # MATRÍCULAS
        # --------------------------------------------------------------

        cabecalhos = [
            "Matrícula",
            "Prefixo",
            "Vínculo",
            "Nome da matrícula",
            "Admissão",
            "Rescisão",
            "Tempo na matrícula",
            "Dias",
            "Situação",
        ]

        for col, titulo in enumerate(
            cabecalhos,
            start=1,
        ):

            c = ws.cell(
                linha_atual,
                col,
                titulo,
            )

            c.font = Font(
                bold=True,
                color=branco,
            )

            c.fill = PatternFill(
                "solid",
                fgColor=azul,
            )

            c.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            c.border = borda

        ws.row_dimensions[
            linha_atual
        ].height = 27

        linha_atual += 1

        for idx, matricula in enumerate(
            pessoa["matriculas"]
        ):

            valores = [
                matricula["matricula"],
                matricula["prefixo"],
                matricula["vinculo"],
                matricula["nome"],
                matricula["admissao"],
                matricula["rescisao"],
                matricula["tempo"],
                matricula["dias"],
                (
                    "ATIVA / PRINCIPAL"
                    if (
                        matricula["principal"]
                        and
                        matricula["rescisao"]
                        == "ATIVO"
                    )
                    else "PRINCIPAL"
                    if matricula["principal"]
                    else ""
                ),
            ]

            for col, valor in enumerate(
                valores,
                start=1,
            ):

                c = ws.cell(
                    linha_atual,
                    col,
                    valor,
                )

                c.border = borda

                c.alignment = Alignment(
                    horizontal=(
                        "left"
                        if col == 4
                        else "center"
                    ),
                    vertical="center",
                    wrap_text=True,
                )

                c.fill = PatternFill(
                    "solid",
                    fgColor=(
                        azul_muito_claro
                        if idx % 2 == 0
                        else branco
                    ),
                )

                if col == 1:
                    c.number_format = "@"

                if matricula["principal"]:

                    c.font = Font(
                        bold=True,
                        color=azul_escuro,
                    )

                if (
                    col == 6
                    and str(valor).upper()
                    == "ATIVO"
                ):

                    c.font = Font(
                        bold=True,
                        color=verde,
                    )

                    c.fill = PatternFill(
                        "solid",
                        fgColor=verde_claro,
                    )

                if (
                    col == 9
                    and "ATIVA"
                    in str(valor).upper()
                ):

                    c.font = Font(
                        bold=True,
                        color=verde,
                    )

                    c.fill = PatternFill(
                        "solid",
                        fgColor=verde_claro,
                    )

            linha_atual += 1

        # --------------------------------------------------------------
        # TOTAL
        # --------------------------------------------------------------

        linha_atual += 1

        ws.merge_cells(
            start_row=linha_atual,
            start_column=1,
            end_row=linha_atual,
            end_column=5,
        )

        c = ws.cell(
            linha_atual,
            1,
            f"TOTAL REAL CONSIDERADO ATÉ "
            f"{data_base.strftime('%d/%m/%Y')}: "
            f"{pessoa['tempo_total']}",
        )

        c.font = Font(
            bold=True,
            color=verde,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=verde_claro,
        )

        c.alignment = Alignment(
            horizontal="right",
            vertical="center",
            wrap_text=True,
        )

        c.border = borda

        for col in range(
            1,
            6,
        ):

            ws.cell(
                linha_atual,
                col,
            ).border = borda

        ws.merge_cells(
            start_row=linha_atual,
            start_column=6,
            end_row=linha_atual,
            end_column=9,
        )

        c = ws.cell(
            linha_atual,
            6,
            f"{pessoa['anos_total']} "
            f"ANOS COMPLETOS",
        )

        c.font = Font(
            size=11,
            bold=True,
            color=verde,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=verde_claro,
        )

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        c.border = borda

        for col in range(
            6,
            10,
        ):

            ws.cell(
                linha_atual,
                col,
            ).border = borda

        linha_atual += 3

    # ==================================================================
    # FORMATAÇÃO
    # ==================================================================

    larguras = {
        1: 16,
        2: 10,
        3: 18,
        4: 42,
        5: 16,
        6: 16,
        7: 26,
        8: 11,
        9: 24,
    }

    for col, largura in (
        larguras.items()
    ):

        ws.column_dimensions[
            get_column_letter(col)
        ].width = largura

    ws.freeze_panes = "A4"

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.page_margins = PageMargins(
        left=0.3,
        right=0.3,
        top=0.5,
        bottom=0.5,
        header=0.2,
        footer=0.2,
    )

    ws.oddFooter.center.text = (
        "Página &P de &N"
    )

    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = "808080"


# ===========================================================================
# ABA DEBUG
# ===========================================================================

def montar_aba_debug(
    workbook,
    mes_alvo,
    ano_alvo,
    data_base,
    data_geracao,
    detalhe_por_pessoa,
    detalhe_geral_por_pessoa,
    dados_por_cpf,
):

    borda = Border(
        left=Side(
            style="thin",
            color=cinza_grade,
        ),
        right=Side(
            style="thin",
            color=cinza_grade,
        ),
        top=Side(
            style="thin",
            color=cinza_grade,
        ),
        bottom=Side(
            style="thin",
            color=cinza_grade,
        ),
    )

    ws = workbook.create_sheet(
        "Debug"
    )

    ws.sheet_view.showGridLines = False

    linha = 1

    # ==================================================================
    # TÍTULO
    # ==================================================================

    ws.merge_cells(
        start_row=linha,
        start_column=1,
        end_row=linha,
        end_column=12,
    )

    c = ws.cell(
        linha,
        1,
        "DEBUG — MEMÓRIA DE CÁLCULO "
        "DOS QUINQUÊNIOS",
    )

    c.font = Font(
        size=16,
        bold=True,
        color=branco,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=azul_escuro,
    )

    c.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    ws.row_dimensions[
        linha
    ].height = 32

    linha += 2

    # ==================================================================
    # 1 — PARÂMETROS
    # ==================================================================

    ws.merge_cells(
        start_row=linha,
        start_column=1,
        end_row=linha,
        end_column=12,
    )

    c = ws.cell(
        linha,
        1,
        "1. PARÂMETROS DA EXECUÇÃO",
    )

    c.font = Font(
        size=12,
        bold=True,
        color=branco,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=azul,
    )

    linha += 1

    parametros = [
        (
            "Mês analisado",
            str(mes_alvo),
        ),
        (
            "Ano analisado",
            str(ano_alvo),
        ),
        (
            "Data-base calculada pelo Python",
            data_base.strftime(
                "%d/%m/%Y"
            ),
        ),
        (
            "Data de geração",
            data_geracao.strftime(
                "%d/%m/%Y"
            ),
        ),
        (
            "Data-base equivalente no Excel",
            formula_excel_como_texto(
                '=EOMONTH(DATE(B5,B4,1),0)'
            ),
        ),
    ]

    for titulo, valor in parametros:

        ws.cell(
            linha,
            1,
            titulo,
        )

        ws.cell(
            linha,
            2,
            valor,
        )

        ws.merge_cells(
            start_row=linha,
            start_column=2,
            end_row=linha,
            end_column=6,
        )

        for col in range(
            1,
            7,
        ):

            c = ws.cell(
                linha,
                col,
            )

            c.border = borda

            c.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            c.fill = PatternFill(
                "solid",
                fgColor=(
                    azul_claro
                    if col == 1
                    else branco
                ),
            )

            if col == 1:

                c.font = Font(
                    bold=True,
                    color=azul_escuro,
                )

        ws.cell(
            linha,
            2,
        ).number_format = "@"

        linha += 1

    linha += 1

    # ==================================================================
    # 2 — CPF
    # ==================================================================

    ws.merge_cells(
        start_row=linha,
        start_column=1,
        end_row=linha,
        end_column=12,
    )

    c = ws.cell(
        linha,
        1,
        "2. TRATAMENTO DO CPF — REGRA APLICADA",
    )

    c.font = Font(
        size=12,
        bold=True,
        color=branco,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=azul_escuro,
    )

    linha += 1

    ws.cell(
        linha,
        1,
        "Situação",
    )

    ws.cell(
        linha,
        2,
        "Python",
    )

    ws.merge_cells(
        start_row=linha,
        start_column=3,
        end_row=linha,
        end_column=8,
    )

    ws.cell(
        linha,
        3,
        "Equivalente / documentação Excel",
    )

    ws.merge_cells(
        start_row=linha,
        start_column=9,
        end_row=linha,
        end_column=12,
    )

    ws.cell(
        linha,
        9,
        "Resultado",
    )

    for col in range(
        1,
        13,
    ):

        c = ws.cell(
            linha,
            col,
        )

        c.font = Font(
            bold=True,
            color=branco,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=azul,
        )

        c.border = borda

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    linha += 1

    regras_cpf = [
        (
            "11 dígitos",
            "Mantém os 11 dígitos exatamente.",
            '=IF(LEN(CPF)=11,CPF,"")',
            "ACEITO — CPF normal",
        ),
        (
            "Vazio / sem dígitos",
            "Não entra no cálculo e vai para a aba Inválidos.",
            '=IF(CPF="","REJEITADO","ACEITO")',
            "REJEITADO",
        ),
        (
            "Quantidade diferente de 11",
            "Não completa, não corta e não corrige automaticamente.",
            '=IF(LEN(CPF)=11,"ACEITO","REJEITADO")',
            "REJEITADO",
        ),
    ]

    for situacao, python_desc, excel_desc, resultado in (
        regras_cpf
    ):

        ws.cell(
            linha,
            1,
            situacao,
        )

        ws.cell(
            linha,
            2,
            python_desc,
        )

        ws.merge_cells(
            start_row=linha,
            start_column=3,
            end_row=linha,
            end_column=8,
        )

        ws.cell(
            linha,
            3,
            formula_excel_como_texto(
                excel_desc
            ),
        )

        ws.merge_cells(
            start_row=linha,
            start_column=9,
            end_row=linha,
            end_column=12,
        )

        ws.cell(
            linha,
            9,
            resultado,
        )

        for col in range(
            1,
            13,
        ):

            c = ws.cell(
                linha,
                col,
            )

            c.border = borda

            c.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            c.fill = PatternFill(
                "solid",
                fgColor=(
                    azul_muito_claro
                    if linha % 2 == 0
                    else branco
                ),
            )

        ws.cell(
            linha,
            3,
        ).number_format = "@"

        linha += 1

    linha += 1

    # ==================================================================
    # 3 — CPF x MATRÍCULA
    # ==================================================================

    ws.merge_cells(
        start_row=linha,
        start_column=1,
        end_row=linha,
        end_column=12,
    )

    c = ws.cell(
        linha,
        1,
        "3. CPF x MATRÍCULA x PREFIXO",
    )

    c.font = Font(
        size=12,
        bold=True,
        color=branco,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=verde,
    )

    linha += 1

    relacoes = [
        (
            "CPF",
            "Chave de agrupamento da pessoa.",
            "Mesmo CPF normalizado = mesmo histórico.",
        ),
        (
            "Matrícula",
            "Identifica o vínculo individual.",
            "Cada matrícula preserva suas próprias datas.",
        ),
        (
            "Prefixo",
            "São os dois primeiros caracteres da matrícula.",
            "01 = ENTITY A | 02 = ENTITY B | 04 = ENTITY C.",
        ),
        (
            "CPF ≠ vínculo",
            "CPF não define o vínculo (ENTITY A/B/C).",
            "Não usar LEFT(CPF,2) para descobrir vínculo.",
        ),
        (
            "Vínculo (fonte)",
            "Tem prioridade quando preenchido.",
            "Prefixo da matrícula é fallback.",
        ),
        (
            "Tempo",
            "Consolidado entre matrículas do mesmo CPF.",
            "Evita contar períodos sobrepostos em dobro.",
        ),
    ]

    ws.cell(
        linha,
        1,
        "Elemento",
    )

    ws.cell(
        linha,
        2,
        "Python",
    )

    ws.merge_cells(
        start_row=linha,
        start_column=3,
        end_row=linha,
        end_column=12,
    )

    ws.cell(
        linha,
        3,
        "Interpretação",
    )

    for col in range(
        1,
        13,
    ):

        c = ws.cell(
            linha,
            col,
        )

        c.font = Font(
            bold=True,
            color=branco,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=azul,
        )

        c.border = borda

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    linha += 1

    for elemento, python_desc, excel_desc in relacoes:

        ws.cell(
            linha,
            1,
            elemento,
        )

        ws.cell(
            linha,
            2,
            python_desc,
        )

        ws.merge_cells(
            start_row=linha,
            start_column=3,
            end_row=linha,
            end_column=12,
        )

        ws.cell(
            linha,
            3,
            excel_desc,
        )

        for col in range(
            1,
            13,
        ):

            c = ws.cell(
                linha,
                col,
            )

            c.border = borda

            c.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            c.fill = PatternFill(
                "solid",
                fgColor=(
                    verde_claro
                    if elemento in (
                        "CPF",
                        "CPF ≠ vínculo",
                    )
                    else (
                        azul_muito_claro
                        if linha % 2 == 0
                        else branco
                    )
                ),
            )

        linha += 1

    linha += 1

    # ==================================================================
    # 4 — PREFIXOS
    # ==================================================================

    ws.merge_cells(
        start_row=linha,
        start_column=1,
        end_row=linha,
        end_column=12,
    )

    c = ws.cell(
        linha,
        1,
        "4. MAPA DE PREFIXOS DA MATRÍCULA",
    )

    c.font = Font(
        size=12,
        bold=True,
        color=branco,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=azul,
    )

    linha += 1

    ws.cell(
        linha,
        1,
        "Prefixo",
    )

    ws.cell(
        linha,
        2,
        "Vínculo Python",
    )

    ws.merge_cells(
        start_row=linha,
        start_column=3,
        end_row=linha,
        end_column=8,
    )

    ws.cell(
        linha,
        3,
        "Equivalente Excel",
    )

    ws.merge_cells(
        start_row=linha,
        start_column=9,
        end_row=linha,
        end_column=12,
    )

    ws.cell(
        linha,
        9,
        "Origem",
    )

    for col in range(
        1,
        13,
    ):

        c = ws.cell(
            linha,
            col,
        )

        c.font = Font(
            bold=True,
            color=branco,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=azul_escuro,
        )

        c.border = borda

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    linha += 1

    for prefixo, nome in (
        ("01", "ENTITY A"),
        ("02", "ENTITY B"),
        ("04", "ENTITY C"),
    ):

        ws.cell(
            linha,
            1,
            prefixo,
        )

        ws.cell(
            linha,
            2,
            nome,
        )

        ws.merge_cells(
            start_row=linha,
            start_column=3,
            end_row=linha,
            end_column=8,
        )

        ws.cell(
            linha,
            3,
            formula_excel_como_texto(
                f'=IF(LEFT(Matricula,2)="{prefixo}","{nome}","")'
            ),
        )

        ws.merge_cells(
            start_row=linha,
            start_column=9,
            end_row=linha,
            end_column=12,
        )

        ws.cell(
            linha,
            9,
            "Prefixo da MATRÍCULA; "
            "nunca do CPF.",
        )

        for col in range(
            1,
            13,
        ):

            c = ws.cell(
                linha,
                col,
            )

            c.border = borda

            c.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            c.fill = PatternFill(
                "solid",
                fgColor=azul_muito_claro,
            )

        ws.cell(
            linha,
            3,
        ).number_format = "@"

        linha += 1

    linha += 1

    # ==================================================================
    # 5 — FLUXO
    # ==================================================================

    ws.merge_cells(
        start_row=linha,
        start_column=1,
        end_row=linha,
        end_column=12,
    )

    c = ws.cell(
        linha,
        1,
        "5. FLUXO DO CÁLCULO",
    )

    c.font = Font(
        size=12,
        bold=True,
        color=branco,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=azul_escuro,
    )

    linha += 1

    fluxo = [
        (
            "1",
            "A consulta retorna as matrículas.",
            "Cada matrícula possui um CPF.",
        ),
        (
            "2",
            "CPF é normalizado.",
            "somente 11 dígitos são aceitos; demais formatos são rejeitados.",
        ),
        (
            "3",
            "CPF normalizado vira a chave.",
            "registros_por_cpf[CPF]",
        ),
        (
            "4",
            "Matrículas do mesmo CPF são reunidas.",
            "Histórico completo da pessoa.",
        ),
        (
            "5",
            "Cada matrícula vira um intervalo.",
            "Admissão = início; rescisão = fim exclusivo.",
        ),
        (
            "6",
            "Intervalos são mesclados.",
            "Sobreposição não gera tempo em dobro.",
        ),
        (
            "7",
            "Hiatos ficam fora.",
            "Cursor virtual não avança.",
        ),
        (
            "8",
            "Quinquênios são calculados.",
            "5, 10, 15, 20... no histórico consolidado.",
        ),
        (
            "9",
            "Matrícula principal é escolhida.",
            "Ativa mais nova; sem ativa, mais nova pela admissão.",
        ),
    ]

    ws.cell(
        linha,
        1,
        "Etapa",
    )

    ws.cell(
        linha,
        2,
        "Python",
    )

    ws.merge_cells(
        start_row=linha,
        start_column=3,
        end_row=linha,
        end_column=12,
    )

    ws.cell(
        linha,
        3,
        "Interpretação",
    )

    for col in range(
        1,
        13,
    ):

        c = ws.cell(
            linha,
            col,
        )

        c.font = Font(
            bold=True,
            color=branco,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=azul,
        )

        c.border = borda

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    linha += 1

    for etapa, python_desc, interpretacao in fluxo:

        ws.cell(
            linha,
            1,
            etapa,
        )

        ws.cell(
            linha,
            2,
            python_desc,
        )

        ws.merge_cells(
            start_row=linha,
            start_column=3,
            end_row=linha,
            end_column=12,
        )

        ws.cell(
            linha,
            3,
            interpretacao,
        )

        for col in range(
            1,
            13,
        ):

            c = ws.cell(
                linha,
                col,
            )

            c.border = borda

            c.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            c.fill = PatternFill(
                "solid",
                fgColor=(
                    azul_muito_claro
                    if linha % 2 == 0
                    else branco
                ),
            )

        linha += 1

    linha += 1

    # ==================================================================
    # 6 — PESSOAS QUE ATINGEM QUINQUÊNIO NO PERÍODO ANALISADO
    # ==================================================================

    ws.merge_cells(
        start_row=linha,
        start_column=1,
        end_row=linha,
        end_column=12,
    )

    c = ws.cell(
        linha,
        1,
        f"6. QUEM ESTÁ TENDO QUINQUÊNIO EM "
        f"{mes_alvo:02d}/{ano_alvo}",
    )

    c.font = Font(
        size=12,
        bold=True,
        color=branco,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=verde,
    )

    linha += 1

    cab_q = [
        "CPF normalizado",
        "CPF bruto observado",
        "Situação do CPF",
        "Nome",
        "Matrícula principal",
        "Prefixo",
        "Vínculo",
        "Quinquênio",
        "Data do quinquênio",
        "Qtd. matrículas",
        "Tempo total",
        "Anos completos",
    ]

    for col, titulo in enumerate(
        cab_q,
        start=1,
    ):

        c = ws.cell(
            linha,
            col,
            titulo,
        )

        c.font = Font(
            bold=True,
            color=branco,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=azul_escuro,
        )

        c.border = borda

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    linha += 1

    total_quinquenios = 0

    if not detalhe_por_pessoa:

        ws.merge_cells(
            start_row=linha,
            start_column=1,
            end_row=linha,
            end_column=12,
        )

        c = ws.cell(
            linha,
            1,
            "Nenhum CPF atingiu quinquênio "
            "no mês analisado.",
        )

        c.fill = PatternFill(
            "solid",
            fgColor=amarelo_claro,
        )

        c.font = Font(
            bold=True,
            color="7F5000",
        )

        c.border = borda

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        linha += 2

    else:

        for pessoa in detalhe_por_pessoa:

            cpf = pessoa["cpf"]
            principal = pessoa["principal"]

            registros = dados_por_cpf[
                cpf
            ][
                "registros_ordenados"
            ]

            cpf_originais = []
            situacoes_cpf = []

            for r in registros:

                if (
                    r["cpf_original"]
                    not in cpf_originais
                ):

                    cpf_originais.append(
                        r["cpf_original"]
                    )

                if (
                    r["cpf_situacao"]
                    not in situacoes_cpf
                ):

                    situacoes_cpf.append(
                        r["cpf_situacao"]
                    )

            cpf_bruto = " | ".join(
                x
                for x in cpf_originais
                if x
            )

            situacao_cpf = " | ".join(
                situacoes_cpf
            )

            for (
                multiplo,
                data_quinquenio,
            ) in pessoa["quinquenios"]:

                total_quinquenios += 1

                valores = [
                    cpf,
                    cpf_bruto,
                    situacao_cpf,
                    pessoa["nome"],
                    principal["matricula"],
                    principal["prefixo"],
                    principal["vinculo"],
                    f"{multiplo} anos",
                    data_quinquenio.strftime(
                        "%d/%m/%Y"
                    ),
                    len(
                        pessoa["matriculas"]
                    ),
                    pessoa["tempo_total"],
                    pessoa["anos_total"],
                ]

                for col, valor in enumerate(
                    valores,
                    start=1,
                ):

                    c = ws.cell(
                        linha,
                        col,
                        valor,
                    )

                    c.border = borda

                    c.alignment = Alignment(
                        horizontal=(
                            "left"
                            if col in (
                                2,
                                3,
                                4,
                                11,
                            )
                            else "center"
                        ),
                        vertical="center",
                        wrap_text=True,
                    )

                    c.fill = PatternFill(
                        "solid",
                        fgColor=verde_claro,
                    )

                    if col == 1:

                        c.number_format = "@"

                    if col == 4:

                        c.font = Font(
                            bold=True,
                            color=azul_escuro,
                        )

                    if col in (
                        8,
                        9,
                    ):

                        c.font = Font(
                            bold=True,
                            color="7F5000",
                        )

                    if col == 9:

                        c.fill = PatternFill(
                            "solid",
                            fgColor=amarelo_claro,
                        )

                linha += 1

    linha += 1

    ws.merge_cells(
        start_row=linha,
        start_column=1,
        end_row=linha,
        end_column=4,
    )

    c = ws.cell(
        linha,
        1,
        f"TOTAL DE QUINQUÊNIOS NO MÊS: "
        f"{total_quinquenios}",
    )

    c.font = Font(
        size=11,
        bold=True,
        color=verde,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=verde_claro,
    )

    c.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    c.border = borda

    for col in range(
        1,
        5,
    ):

        ws.cell(
            linha,
            col,
        ).border = borda

    ws.merge_cells(
        start_row=linha,
        start_column=5,
        end_row=linha,
        end_column=12,
    )

    c = ws.cell(
        linha,
        5,
        "O resultado acima foi calculado "
        "sobre o histórico consolidado de cada CPF.",
    )

    c.font = Font(
        bold=True,
        color=azul_escuro,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=azul_claro,
    )

    c.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    c.border = borda

    for col in range(
        5,
        13,
    ):

        ws.cell(
            linha,
            col,
        ).border = borda

    linha += 2

    # ==================================================================
    # 7 — CPF REAIS PROCESSADOS
    # ==================================================================

    ws.merge_cells(
        start_row=linha,
        start_column=1,
        end_row=linha,
        end_column=12,
    )

    c = ws.cell(
        linha,
        1,
        "7. CPF PROCESSADOS — AUDITORIA",
    )

    c.font = Font(
        size=12,
        bold=True,
        color=branco,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=azul,
    )

    linha += 1

    cab_auditoria = [
        "CPF normalizado",
        "CPF bruto",
        "Situação",
        "Nome",
        "Qtd. matrículas",
        "Matrícula principal",
        "Prefixo",
        "Vínculo",
        "Tempo Total CPF",
        "Anos completos",
        "Quinquênio no mês?",
        "Datas dos quinquênios",
    ]

    for col, titulo in enumerate(
        cab_auditoria,
        start=1,
    ):

        c = ws.cell(
            linha,
            col,
            titulo,
        )

        c.font = Font(
            bold=True,
            color=branco,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=azul_escuro,
        )

        c.border = borda

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    linha += 1

    for pessoa in detalhe_geral_por_pessoa:

        cpf = pessoa["cpf"]

        info = dados_por_cpf[
            cpf
        ]

        registros = info[
            "registros_ordenados"
        ]

        cpf_originais = []
        situacoes = []

        for r in registros:

            if (
                r["cpf_original"]
                not in cpf_originais
            ):

                cpf_originais.append(
                    r["cpf_original"]
                )

            if (
                r["cpf_situacao"]
                not in situacoes
            ):

                situacoes.append(
                    r["cpf_situacao"]
                )

        quinquenios_txt = ", ".join(
            f"{m} anos — "
            f"{d.strftime('%d/%m/%Y')}"
            for m, d in info["quinquenios"]
        )

        quinquenio_mes = any(
            d.month == mes_alvo
            and d.year == ano_alvo
            for _, d in info["quinquenios"]
        )

        valores = [
            cpf,
            " | ".join(
                cpf_originais
            ),
            " | ".join(
                situacoes
            ),
            pessoa["nome"],
            len(registros),
            info["principal"]["matricula"],
            info["principal"]["prefixo"],
            info["principal"]["vinculo"],
            info["tempo_total"],
            info["anos_total"],
            "SIM" if quinquenio_mes else "NÃO",
            quinquenios_txt,
        ]

        for col, valor in enumerate(
            valores,
            start=1,
        ):

            c = ws.cell(
                linha,
                col,
                valor,
            )

            c.border = borda

            c.alignment = Alignment(
                horizontal=(
                    "left"
                    if col in (
                        2,
                        3,
                        4,
                        9,
                        12,
                    )
                    else "center"
                ),
                vertical="center",
                wrap_text=True,
            )

            c.fill = PatternFill(
                "solid",
                fgColor=(
                    verde_claro
                    if quinquenio_mes
                    else (
                        azul_muito_claro
                        if linha % 2 == 0
                        else branco
                    )
                ),
            )

            if col == 1:
                c.number_format = "@"

            if col == 11:

                if valor == "SIM":

                    c.font = Font(
                        bold=True,
                        color=verde,
                    )

                    c.fill = PatternFill(
                        "solid",
                        fgColor=verde_claro,
                    )

                else:

                    c.font = Font(
                        color=cinza,
                    )

        linha += 1

    linha += 1

    # ==================================================================
    # 8 — REGRAS DE TEMPO
    # ==================================================================

    ws.merge_cells(
        start_row=linha,
        start_column=1,
        end_row=linha,
        end_column=12,
    )

    c = ws.cell(
        linha,
        1,
        "8. REGRAS DE TEMPO DE SERVIÇO",
    )

    c.font = Font(
        size=12,
        bold=True,
        color=branco,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=azul_escuro,
    )

    linha += 1

    regras_tempo = [
        (
            "Data-base",
            "Último dia real do mês informado.",
            '=EOMONTH(DATE(Ano,Mes,1),0)',
        ),
        (
            "Rescisão",
            "Fim exclusivo. O dia da rescisão não entra como dia trabalhado.",
            '=Rescisao-Admissao',
        ),
        (
            "Continuidade",
            "Nova admissão igual à rescisão anterior não cria hiato.",
            'Nova_Admissao<=Rescisao_Anterior',
        ),
        (
            "Sobreposição",
            "Matrículas simultâneas são fundidas.",
            "UNION / MERGE dos períodos",
        ),
        (
            "Hiato",
            "Período sem vínculo não entra no tempo.",
            "Cursor virtual não avança",
        ),
        (
            "Meses de serviço",
            "Diferença calculada somente por ANO + MÊS. O DIA é ignorado.",
            '=(AnoFim-AnoInicio)*12+(MesFim-MesInicio)',
        ),
        (
            "Anos completos",
            "Obtidos a partir dos meses completos acumulados.",
            '=INT(MesesCompletos/12)',
        ),
        (
            "Quinquênio",
            "60 meses por marco: 60, 120, 180, 240...",
            '=INT(MesesAcumulados/60)',
        ),
        (
            "Exceção do dia",
            "O dia pode ser anterior ao aniversário, desde que esteja no mesmo MÊS/ANO.",
            "15/06/2012 → 14/06/2026 = 168 meses",
        ),
        (
            "Exceção do mês",
            "Estar um mês antes do aniversário NÃO completa o ano.",
            "01/11/2021 → 30/10/2026 = 59 meses",
        ),
        (
            "Data-base rígida",
            "Nenhum quinquênio pode ser projetado para MÊS/ANO posterior à data-base.",
            "Marco <= MÊS/ANO da data-base",
        ),
    ]
    ws.cell(
        linha,
        1,
        "Regra",
    )

    ws.cell(
        linha,
        2,
        "Python",
    )

    ws.merge_cells(
        start_row=linha,
        start_column=3,
        end_row=linha,
        end_column=12,
    )

    ws.cell(
        linha,
        3,
        "Equivalente / documentação Excel",
    )

    for col in range(
        1,
        13,
    ):

        c = ws.cell(
            linha,
            col,
        )

        c.font = Font(
            bold=True,
            color=branco,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=azul,
        )

        c.border = borda

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    linha += 1

    for regra, python_desc, excel_desc in regras_tempo:

        ws.cell(
            linha,
            1,
            regra,
        )

        ws.cell(
            linha,
            2,
            python_desc,
        )

        ws.merge_cells(
            start_row=linha,
            start_column=3,
            end_row=linha,
            end_column=12,
        )

        ws.cell(
            linha,
            3,
            formula_excel_como_texto(
                excel_desc
            ),
        )

        for col in range(
            1,
            13,
        ):

            c = ws.cell(
                linha,
                col,
            )

            c.border = borda

            c.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            c.fill = PatternFill(
                "solid",
                fgColor=(
                    azul_muito_claro
                    if linha % 2 == 0
                    else branco
                ),
            )

        ws.cell(
            linha,
            3,
        ).number_format = "@"

        linha += 1

    linha += 1

    # ==================================================================
    # 9 — CONCLUSÃO
    # ==================================================================

    ws.merge_cells(
        start_row=linha,
        start_column=1,
        end_row=linha,
        end_column=12,
    )

    c = ws.cell(
        linha,
        1,
        "9. CONCLUSÃO DA AUDITORIA",
    )

    c.font = Font(
        size=12,
        bold=True,
        color=branco,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=azul_escuro,
    )

    linha += 1

    conclusoes = [
        "O CPF é armazenado e tratado como texto.",
        "CPF com 11 dígitos permanece exatamente com 11 dígitos.",
        "CPF com quantidade diferente de 11 dígitos é rejeitado.",
        "CPF com quantidade diferente de 11 dígitos é rejeitado.",
        "O CPF normalizado é a chave de agrupamento do histórico da pessoa.",
        "A matrícula é mantida individual dentro do histórico agrupado pelo CPF.",
        "O prefixo usado para vínculo vem da matrícula, nunca do CPF.",
        "01 = ENTITY A; 02 = ENTITY B; 04 = ENTITY C.",
        "O vínculo retornado pela consulta tem prioridade; o prefixo da matrícula é fallback.",
        "Matrículas sobrepostas não contam tempo em dobro.",
        "Hiatos sem vínculo não são contabilizados.",
        "Rescisões são exclusivas.",
        "A matrícula principal é a ativa mais nova; sem ativa, é a mais nova pela admissão.",
        "As fórmulas exibidas no Debug estão gravadas como TEXTO.",
        "O cálculo do quinquênio ignora deliberadamente o DIA.",
        "O cálculo NÃO ignora o MÊS nem o ANO.",
        "Se admissão e fim estiverem no mesmo MÊS/ANO, o ano é considerado completo.",
        "Se o fim estiver no mês anterior ao aniversário, o ano NÃO é completado.",
        "Exemplo: 15/06/2012 → 14/06/2026 = 168 meses = 14 anos.",
        "Exemplo: 01/11/2021 → 30/10/2026 = 59 meses = 4 anos e 11 meses.",
        "Exemplo: 01/11/2021 → 30/11/2026 = 60 meses = 5 anos.",
        "Nenhum quinquênio pode ultrapassar o MÊS/ANO da data-base.",
    ]

    for texto in conclusoes:

        ws.merge_cells(
            start_row=linha,
            start_column=1,
            end_row=linha,
            end_column=12,
        )

        c = ws.cell(
            linha,
            1,
            "• " + texto,
        )

        c.border = borda

        c.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=(
                azul_muito_claro
                if linha % 2 == 0
                else branco
            ),
        )

        c.font = Font(
            color=cinza,
        )

        linha += 1

    # ==================================================================
    # FORMATAÇÃO
    # ==================================================================

    larguras = {
        1: 22,
        2: 30,
        3: 34,
        4: 42,
        5: 22,
        6: 18,
        7: 18,
        8: 18,
        9: 25,
        10: 18,
        11: 20,
        12: 34,
    }

    for col, largura in (
        larguras.items()
    ):

        ws.column_dimensions[
            get_column_letter(col)
        ].width = largura

    ws.freeze_panes = "A4"

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.page_margins = PageMargins(
        left=0.3,
        right=0.3,
        top=0.5,
        bottom=0.5,
        header=0.2,
        footer=0.2,
    )

    ws.oddFooter.center.text = (
        "Página &P de &N"
    )

    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = "808080"


# ===========================================================================
# PROCESSAMENTO
# ===========================================================================

def executar_calculo_df(
    df,
    mes_alvo,
    ano_alvo,
    caminho_saida,
    logger=None,
    modo="CPF",
):

    data_geracao = date.today()

    data_base = ultimo_dia_do_mes(
        mes_alvo,
        ano_alvo,
    )

    if logger:

        logger(
            "Data-base do cálculo: "
            f"{data_base.strftime('%d/%m/%Y')}"
        )

        logger(
            f"Normalizando {modo} e matrículas..."
        )

        if str(modo).upper() == "CPF":
            logger(
                "Regra CPF: somente 11 dígitos são aceitos; "
                "demais formatos vão para Inválidos."
            )
        else:
            logger(
                "Regra PIS: 11 dígitos mantidos; "
                "12 dígitos perdem o primeiro; "
                "demais formatos vão para Inválidos."
            )

    (
        registros_por_cpf,
        registros_invalidos,
    ) = preparar_registros(
        df,
        data_base,
        logger,
        modo,
    )

    resultados_finais = []

    cpf_selecionados = set()

    dados_por_cpf = {}

    # ==================================================================
    # PROCESSAMENTO POR CPF
    # ==================================================================

    for cpf, registros in (
        registros_por_cpf.items()
    ):

        registros_ordenados = sorted(
            registros,
            key=lambda r: r["inicio"],
        )

        principal = (
            selecionar_matricula_principal(
                registros_ordenados
            )
        )

        intervalos = [
            {
                "inicio": r["inicio"],
                "fim": r["fim"],
                "aberto": r["aberto"],
            }
            for r in registros_ordenados
        ]

        mesclados = mesclar_intervalos(
            intervalos
        )

        quinquenios = (
            calcular_datas_quinquenio(
                mesclados,
                data_base,
            )
        )

        tempo_total = formatar_tempo(
            mesclados,
            data_base,
        )

        anos_total = anos_completos(
            mesclados,
            data_base,
        )

        dados_por_cpf[cpf] = {
            "registros_ordenados":
                registros_ordenados,

            "mesclados":
                mesclados,

            "principal":
                principal,

            "tempo_total":
                tempo_total,

            "anos_total":
                anos_total,

            "quinquenios":
                quinquenios,
        }

        for (
            multiplo,
            data_atingida,
        ) in quinquenios:

            if (
                data_atingida.month
                == mes_alvo
                and
                data_atingida.year
                == ano_alvo
            ):

                cpf_selecionados.add(
                    cpf
                )

                resultados_finais.append(
                    {
                        "Matrícula":
                            principal["matricula"],

                        "Vínculo":
                            principal["vinculo"],

                        "Nome":
                            principal["nome"],

                        "Função":
                            principal["funcao"],

                        "Lotação":
                            principal["lotacao"],

                        "Admissão":
                            principal["inicio"].strftime(
                                "%d/%m/%Y"
                            ),

                        "Mês":
                            [
                                "JANEIRO",
                                "FEVEREIRO",
                                "MARÇO",
                                "ABRIL",
                                "MAIO",
                                "JUNHO",
                                "JULHO",
                                "AGOSTO",
                                "SETEMBRO",
                                "OUTUBRO",
                                "NOVEMBRO",
                                "DEZEMBRO",
                            ][mes_alvo - 1],

                        "Data-base":
                            data_base,

                        "Quinquênio":
                            multiplo,
                    }
                )

    # ==================================================================
    # DETALHE QUINQUÊNIOS
    # ==================================================================

    detalhe_por_pessoa = []

    for cpf in sorted(
        cpf_selecionados,
        key=lambda p:
        str(
            dados_por_cpf[p]["principal"]["nome"]
        ).upper(),
    ):

        info = dados_por_cpf[cpf]

        registros_ordenados = (
            info["registros_ordenados"]
        )

        principal = info["principal"]

        matriculas_pessoa = []

        for r in registros_ordenados:

            fim_efetivo = (
                data_base
                if r["aberto"]
                else r["fim"]
            )

            if (
                fim_efetivo is not None
                and fim_efetivo > data_base
            ):

                fim_efetivo = data_base

            dias = dias_intervalo(
                r["inicio"],
                fim_efetivo,
            )

            tempo_individual = (
                formatar_tempo(
                    [
                        {
                            "inicio":
                                r["inicio"],
                            "fim":
                                r["fim"],
                            "aberto":
                                r["aberto"],
                        }
                    ],
                    data_base,
                )
            )

            matriculas_pessoa.append(
                {
                    "matricula":
                        r["matricula"],

                    "prefixo":
                        r["prefixo"],

                    "vinculo":
                        r["vinculo"],

                    "nome":
                        r["nome"],

                    "admissao":
                        r["inicio"].strftime(
                            "%d/%m/%Y"
                        ),

                    "rescisao":
                        (
                            "ATIVO"
                            if r["aberto"]
                            else r["fim"].strftime(
                                "%d/%m/%Y"
                            )
                        ),

                    "tempo":
                        tempo_individual,

                    "dias":
                        dias,

                    "lotacao":
                        r["lotacao"],

                    "jornada":
                        r["jornada"],

                    "principal":
                        r is principal,

                    "cpf_original":
                        r["cpf_original"],

                    "cpf_situacao":
                        r["cpf_situacao"],
                }
            )

        detalhe_por_pessoa.append(
            {
                "nome":
                    principal["nome"],

                "cpf":
                    cpf,

                "principal":
                    principal,

                "quinquenios":
                    [
                        (m, d)
                        for m, d
                        in info["quinquenios"]
                        if (
                            d.month == mes_alvo
                            and
                            d.year == ano_alvo
                        )
                    ],

                "matriculas":
                    matriculas_pessoa,

                "tempo_total":
                    info["tempo_total"],

                "anos_total":
                    info["anos_total"],
            }
        )

    # ==================================================================
    # DETALHE GERAL
    # ==================================================================

    detalhe_geral_por_pessoa = []

    for cpf in sorted(
        dados_por_cpf.keys(),
        key=lambda p:
        str(
            dados_por_cpf[p]["principal"]["nome"]
        ).upper(),
    ):

        info = dados_por_cpf[cpf]

        registros_ordenados = (
            info["registros_ordenados"]
        )

        principal = info["principal"]

        matriculas_geral = []

        for r in registros_ordenados:

            fim_efetivo = (
                data_base
                if r["aberto"]
                else r["fim"]
            )

            if (
                fim_efetivo is not None
                and fim_efetivo > data_base
            ):

                fim_efetivo = data_base

            dias = dias_intervalo(
                r["inicio"],
                fim_efetivo,
            )

            tempo_individual = (
                formatar_tempo(
                    [
                        {
                            "inicio":
                                r["inicio"],
                            "fim":
                                r["fim"],
                            "aberto":
                                r["aberto"],
                        }
                    ],
                    data_base,
                )
            )

            matriculas_geral.append(
                {
                    "matricula":
                        r["matricula"],

                    "prefixo":
                        r["prefixo"],

                    "vinculo":
                        r["vinculo"],

                    "nome":
                        r["nome"],

                    "admissao":
                        r["inicio"].strftime(
                            "%d/%m/%Y"
                        ),

                    "rescisao":
                        (
                            "ATIVO"
                            if r["aberto"]
                            else r["fim"].strftime(
                                "%d/%m/%Y"
                            )
                        ),

                    "tempo":
                        tempo_individual,

                    "dias":
                        dias,

                    "lotacao":
                        r["lotacao"],

                    "jornada":
                        r["jornada"],

                    "principal":
                        r is principal,

                    "cpf_original":
                        r["cpf_original"],

                    "cpf_situacao":
                        r["cpf_situacao"],
                }
            )

        quinquenios_mes = [
            (m, d)
            for m, d in info["quinquenios"]
            if (
                d.month == mes_alvo
                and
                d.year == ano_alvo
            )
        ]

        detalhe_geral_por_pessoa.append(
            {
                "nome":
                    principal["nome"],

                "cpf":
                    cpf,

                "principal":
                    principal,

                "quinquenios_mes":
                    quinquenios_mes,

                "matriculas":
                    matriculas_geral,

                "tempo_total":
                    info["tempo_total"],

                "anos_total":
                    info["anos_total"],
            }
        )

    # ==================================================================
    # RESULTADO
    # ==================================================================

    if not resultados_finais:

        raise ValueError(
            f"Nenhum funcionário completa "
            f"quinquênio em "
            f"{mes_alvo:02d}/{ano_alvo}."
        )

    df_saida = (
        pd.DataFrame(
            resultados_finais
        )
        .sort_values(
            "Nome"
        )
    )

    if logger:

        logger(
            "Funcionários encontrados no mês: "
            f"{len(df_saida):,}."
        )

        logger(
            "Pessoas analisadas: "
            f"{len(detalhe_geral_por_pessoa):,}."
        )

        logger(
            "Registros inválidos separados: "
            f"{len(registros_invalidos):,}."
        )

        logger(
            "Gerando Excel..."
        )

    # ==================================================================
    # EXCEL
    # ==================================================================

    with pd.ExcelWriter(
        caminho_saida,
        engine="openpyxl",
        datetime_format="DD/MM/YYYY",
    ) as writer:

        # --------------------------------------------------------------
        # PRINCIPAL
        # --------------------------------------------------------------

        df_saida.to_excel(
            writer,
            index=False,
            sheet_name="Quinquenios",
        )

        workbook = writer.book

        borda = Border(
            left=Side(
                style="thin",
                color=cinza_grade,
            ),
            right=Side(
                style="thin",
                color=cinza_grade,
            ),
            top=Side(
                style="thin",
                color=cinza_grade,
            ),
            bottom=Side(
                style="thin",
                color=cinza_grade,
            ),
        )

        planilha = writer.sheets[
            "Quinquenios"
        ]

        # ==============================================================
        # FORMATAÇÃO DA ABA 'Quinquenios' CONFORME MODELO
        # ==============================================================

        planilha.sheet_view.showGridLines = False

        borda_modelo = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # --------------------------------------------------------------
        # CABEÇALHO
        # --------------------------------------------------------------

        estilos_cabecalho = {
            1: {"fill": "D9D9D9", "align": Alignment()},
            2: {"fill": "D9D9D9", "align": Alignment()},
            3: {"fill": "D9D9D9", "align": Alignment()},
            4: {"fill": "D9D9D9", "align": Alignment()},
            5: {"fill": "D9D9D9", "align": Alignment()},
            6: {"fill": "D9D9D9", "align": Alignment()},
            7: {"fill": "92D050", "align": Alignment()},
            8: {"fill": "FCD5B5", "align": Alignment()},
            9: {
                "fill": "D9D9D9",
                "align": Alignment(horizontal="center"),
            },
        }

        for col_idx, cfg in estilos_cabecalho.items():

            cel = planilha.cell(1, col_idx)

            cel.font = Font(
                name="Arial",
                size=10,
                bold=True,
                color="000000",
            )

            cel.fill = PatternFill(
                "solid",
                fgColor=cfg["fill"],
            )

            cel.border = borda_modelo
            cel.alignment = cfg["align"]

            if col_idx == 8:
                cel.number_format = "mm-dd-yy"
            elif col_idx == 9:
                cel.number_format = "0"
            else:
                cel.number_format = "General"

        # --------------------------------------------------------------
        # DADOS
        # --------------------------------------------------------------

        ultima_linha = len(df_saida) + 1

        for row_idx in range(2, ultima_linha + 1):

            # A:F — padrão neutro do modelo
            for col_idx in range(1, 7):

                cel = planilha.cell(
                    row_idx,
                    col_idx,
                )

                cel.font = Font(
                    name="Arial",
                    size=10,
                    color="000000",
                )

                cel.fill = PatternFill(
                    fill_type=None,
                )

                cel.border = borda_modelo
                cel.alignment = Alignment()
                cel.number_format = "General"

            # G — Mês
            cel = planilha.cell(
                row_idx,
                7,
            )

            cel.font = Font(
                name="Arial",
                size=10,
                bold=True,
                color="000000",
            )

            cel.fill = PatternFill(
                "solid",
                fgColor="92D050",
            )

            cel.border = borda_modelo
            cel.alignment = Alignment(
                vertical="center",
            )
            cel.number_format = r"dd/mm/yyyy\ "

            # H — Data-base
            cel = planilha.cell(
                row_idx,
                8,
            )

            cel.font = Font(
                name="Arial",
                size=10,
                bold=True,
                color="000000",
            )

            cel.fill = PatternFill(
                "solid",
                fgColor="F0A800",
            )

            cel.border = borda_modelo
            cel.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            cel.number_format = "mm-dd-yy"

            # I — Quinquênio
            cel = planilha.cell(
                row_idx,
                9,
            )

            cel.font = Font(
                name="Arial",
                size=10,
                color="000000",
            )

            cel.fill = PatternFill(
                "solid",
                fgColor="BFBFBF",
            )

            cel.border = borda_modelo
            cel.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            cel.number_format = "0"

        # Matrícula permanece textual, como no modelo.
        for row_idx in range(2, ultima_linha + 1):
            planilha.cell(row_idx, 1).number_format = "@"

        # --------------------------------------------------------------
        # DIMENSÕES / FILTRO / CONGELAMENTO
        # --------------------------------------------------------------

        larguras_principal = {
            "A": 11.7109375,
            "B": 16.5703125,
            "C": 40.5703125,
            "D": 55.42578125,
            "E": 47.0,
            "F": 12.0,
            "G": 10.140625,
            "H": 12.28515625,
            "I": 16.140625,
        }

        for coluna, largura in larguras_principal.items():
            planilha.column_dimensions[coluna].width = largura

        planilha.auto_filter.ref = f"A1:I{ultima_linha}"
        planilha.freeze_panes = "A2"

        # O modelo usa estas mesmas configurações de impressão.
        planilha.sheet_properties.pageSetUpPr.fitToPage = True
        planilha.page_setup.orientation = "landscape"
        planilha.page_setup.paperSize = planilha.PAPERSIZE_A4
        planilha.page_setup.fitToWidth = None
        planilha.page_setup.fitToHeight = 0

        planilha.page_margins = PageMargins(
            left=0.3,
            right=0.3,
            top=0.5,
            bottom=0.5,
            header=0.2,
            footer=0.2,
        )

        planilha.oddFooter.center.text = "Página &P de &N"
        planilha.oddFooter.center.size = 8
        planilha.oddFooter.center.color = "808080"

        # --------------------------------------------------------------
        # DETALHE POR MATRÍCULA
        # --------------------------------------------------------------

        montar_aba_detalhe_matricula(
            workbook,
            mes_alvo,
            ano_alvo,
            data_base,
            data_geracao,
            detalhe_por_pessoa,
        )

        # --------------------------------------------------------------
        # DETALHE GERAL
        # --------------------------------------------------------------

        montar_aba_detalhe_geral(
            workbook,
            mes_alvo,
            ano_alvo,
            data_base,
            data_geracao,
            detalhe_geral_por_pessoa,
        )

        # --------------------------------------------------------------
        # INVÁLIDOS
        # --------------------------------------------------------------

        montar_aba_invalidos(
            workbook,
            registros_invalidos,
        )

        # --------------------------------------------------------------
        # DEBUG
        # --------------------------------------------------------------

        montar_aba_debug(
            workbook,
            mes_alvo,
            ano_alvo,
            data_base,
            data_geracao,
            detalhe_por_pessoa,
            detalhe_geral_por_pessoa,
            dados_por_cpf,
        )

        ajustar_rotulos_por_modo(
            workbook,
            modo,
        )

        # --------------------------------------------------------------
        # RECÁLCULO EXCEL
        # --------------------------------------------------------------

        try:

            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
            workbook.calculation.calcMode = "auto"

        except Exception:

            pass

    if logger:

        logger(
            f"Concluído: {caminho_saida}"
        )

    return (
        caminho_saida,
        len(df_saida),
        len(detalhe_geral_por_pessoa),
    )


# ===========================================================================
# AJUSTE DE RÓTULOS CPF/PIS
# ===========================================================================

def ajustar_rotulos_por_modo(
    workbook,
    modo,
):
    """
    Quando o modo PIS estiver selecionado, troca os rótulos de auditoria e
    detalhes que originalmente usam "CPF". Os valores das chaves continuam
    sendo os calculados pelo modo escolhido.
    """

    if str(modo).upper() != "PIS":
        return

    for nome_planilha in (
        "Inválidos",
        "Detalhe por Mat. Geral",
        "Detalhe por Matricula",
        "Debug",
    ):

        if nome_planilha not in workbook.sheetnames:
            continue

        ws = workbook[nome_planilha]

        for row in ws.iter_rows():

            for cell in row:

                if not isinstance(cell.value, str):
                    continue

                cell.value = (
                    cell.value
                    .replace("CPF", "PIS")
                    .replace("cpf", "pis")
                )


# ===========================================================================
# GUI
# ===========================================================================

def centralizar(win):

    win.update_idletasks()

    w = win.winfo_width()
    h = win.winfo_height()

    sw = win.winfo_screenwidth()
    sh = win.winfo_screenheight()

    x = max(
        0,
        (sw - w) // 2,
    )

    y = max(
        0,
        (sh - h) // 2,
    )

    win.geometry(
        f"{w}x{h}+{x}+{y}"
    )


def montar_gui():

    root = tk.Tk()

    root.title(
        "Quintennial Data Pipeline — Gerador de Quinquênios"
    )

    root.geometry(
        "900x640"
    )

    root.minsize(
        820,
        580,
    )

    root.configure(
        bg=BG_MAIN
    )

    style = ttk.Style(
        root
    )

    try:

        style.theme_use(
            "clam"
        )

    except tk.TclError:
        pass

    style.configure(
        "TFrame",
        background=BG_MAIN,
    )

    style.configure(
        "Card.TFrame",
        background=BG_CARD,
    )

    style.configure(
        "TLabel",
        background=BG_CARD,
        foreground=TEXT_MAIN,
        font=(FONT, 10),
    )

    style.configure(
        "Muted.TLabel",
        background=BG_CARD,
        foreground=TEXT_MUTED,
        font=(FONT, 9),
    )

    style.configure(
        "Title.TLabel",
        background=BG_MAIN,
        foreground=TEXT_MAIN,
        font=(FONT, 18, "bold"),
    )

    style.configure(
        "Header.TLabel",
        background=BG_MAIN,
        foreground=TEXT_MUTED,
        font=(FONT, 10),
    )

    style.configure(
        "TButton",
        font=(FONT, 9),
        padding=(8, 5),
    )

    style.configure(
        "Secondary.TButton",
        font=(FONT, 9),
        padding=(8, 5),
    )

    style.configure(
        "Accent.TButton",
        font=(FONT, 9, "bold"),
        padding=(10, 6),
        foreground="#FFFFFF",
        background=ACCENT_DARK,
    )

    style.map(
        "TButton",
        background=[
            ("active", ACCENT_HOVER)
        ],
    )

    style.map(
        "Secondary.TButton",
        background=[
            ("active", BG_SOFT)
        ],
    )

    style.map(
        "Accent.TButton",
        background=[
            ("active", ACCENT_HOVER),
            ("disabled", "#B9C8D6"),
        ],
        foreground=[
            ("disabled", "#FFFFFF"),
        ],
    )

    # ==================================================================
    # VARIÁVEIS
    # ==================================================================

    mes_var = tk.StringVar(
        value=f"{date.today().month:02d}"
    )

    ano_var = tk.StringVar(
        value=str(
            date.today().year
        )
    )

    pasta_var = tk.StringVar(
        value=PASTA_SAIDA_PADRAO
    )

    status_var = tk.StringVar(
        value="Pronto."
    )

    modo_var = tk.StringVar(
        value="CPF"
    )

    cabecalho_var = tk.StringVar(
        value="Consulta SQL → agrupamento por CPF → cálculo → relatório Excel"
    )

    root.grid_columnconfigure(
        0,
        weight=1,
    )

    root.grid_rowconfigure(
        2,
        weight=1,
    )

    # ==================================================================
    # HEADER
    # ==================================================================

    header = ttk.Frame(
        root
    )

    header.grid(
        row=0,
        column=0,
        sticky="ew",
        padx=24,
        pady=(18, 6),
    )

    ttk.Label(
        header,
        text="Gerador de Quinquênios",
        style="Title.TLabel",
    ).pack(
        anchor="w"
    )

    ttk.Label(
        header,
        textvariable=cabecalho_var,
        style="Header.TLabel",
    ).pack(
        anchor="w",
        pady=(2, 0),
    )

    # ==================================================================
    # CONFIGURAÇÃO
    # ==================================================================

    card = ttk.Frame(
        root,
        style="Card.TFrame",
        padding=16,
    )

    card.grid(
        row=1,
        column=0,
        sticky="ew",
        padx=24,
        pady=8,
    )

    card.grid_columnconfigure(
        0,
        weight=0,
    )

    card.grid_columnconfigure(
        1,
        weight=0,
    )

    card.grid_columnconfigure(
        2,
        weight=0,
    )

    card.grid_columnconfigure(
        3,
        weight=4,
    )

    card.grid_columnconfigure(
        4,
        weight=0,
    )

    card.grid_columnconfigure(
        5,
        weight=3,
    )

    # Labels row 0
    ttk.Label(
        card,
        text="Mês",
    ).grid(
        row=0,
        column=0,
        sticky="w",
    )

    ttk.Label(
        card,
        text="Ano",
    ).grid(
        row=0,
        column=1,
        sticky="w",
        padx=(10, 0),
    )

    ttk.Label(
        card,
        text="Agrupar por",
    ).grid(
        row=0,
        column=2,
        sticky="w",
        padx=(10, 0),
    )

    ttk.Label(
        card,
        text="Pasta de saída",
    ).grid(
        row=0,
        column=3,
        sticky="w",
        padx=(10, 0),
    )

    ttk.Label(
        card,
        text="Nome do arquivo",
    ).grid(
        row=0,
        column=5,
        sticky="w",
        padx=(12, 0),
    )

    # Widgets row 1
    mes_cb = ttk.Combobox(
        card,
        textvariable=mes_var,
        values=[
            f"{i:02d}"
            for i in range(
                1,
                13,
            )
        ],
        width=7,
        state="readonly",
    )

    mes_cb.grid(
        row=1,
        column=0,
        sticky="ew",
        pady=(3, 3),
    )

    ttk.Entry(
        card,
        textvariable=ano_var,
    ).grid(
        row=1,
        column=1,
        sticky="ew",
        padx=(10, 0),
        pady=(3, 3),
    )

    modo_cb = ttk.Combobox(
        card,
        textvariable=modo_var,
        values=["CPF", "PIS"],
        width=10,
        state="readonly",
    )

    modo_cb.grid(
        row=1,
        column=2,
        sticky="ew",
        padx=(10, 0),
        pady=(3, 3),
    )

    ttk.Entry(
        card,
        textvariable=pasta_var,
    ).grid(
        row=1,
        column=3,
        sticky="ew",
        padx=(10, 0),
        pady=(3, 3),
    )

    def escolher_pasta():

        pasta = filedialog.askdirectory(
            title="Pasta de saída",
            initialdir=pasta_var.get(),
        )

        if not pasta:
            return

        pasta_var.set(pasta)

    ttk.Button(
        card,
        text="📂Selecionar Pasta",
        style="Secondary.TButton",
        command=escolher_pasta,
    ).grid(
        row=1,
        column=4,
        sticky="e",
        padx=(8, 0),
        pady=(3, 3),
    )

    nome_var = tk.StringVar(
        value=(
            f"quinquenios_CPF_"
            f"{date.today().month:02d}_"
            f"{date.today().year}.xlsx"
        )
    )

    ttk.Entry(
        card,
        textvariable=nome_var,
        width=28,
    ).grid(
        row=1,
        column=5,
        sticky="ew",
        padx=(12, 0),
        pady=(3, 3),
    )

    def atualizar_nome_arquivo(*_):
        try:
            mes = int(mes_var.get())
            ano = int(ano_var.get())
            modo = modo_var.get().strip().upper() or "CPF"
            nome_var.set(
                f"quinquenios_{modo}_{mes:02d}_{ano}.xlsx"
            )
        except ValueError:
            pass

    def atualizar_cabecalho(*_):
        modo = modo_var.get().strip().upper() or "CPF"
        cabecalho_var.set(
            "Consulta SQL → "
            f"agrupamento por {modo} → "
            "cálculo → relatório Excel"
        )

    mes_var.trace_add("write", atualizar_nome_arquivo)
    ano_var.trace_add("write", atualizar_nome_arquivo)
    modo_var.trace_add("write", atualizar_nome_arquivo)
    modo_var.trace_add("write", atualizar_cabecalho)

    # ==================================================================
    # BODY
    # ==================================================================

    body = ttk.Frame(
        root
    )

    body.grid(
        row=2,
        column=0,
        sticky="nsew",
        padx=24,
        pady=(0, 14),
    )

    body.grid_columnconfigure(
        0,
        weight=1,
    )

    body.grid_rowconfigure(
        0,
        weight=1,
    )

    # ==================================================================
    # LOG
    # ==================================================================

    log_card = ttk.Frame(
        body,
        style="Card.TFrame",
        padding=14,
    )

    log_card.grid(
        row=0,
        column=0,
        sticky="nsew",
    )

    log_card.grid_columnconfigure(
        0,
        weight=1,
    )

    log_card.grid_rowconfigure(
        1,
        weight=1,
    )

    ttk.Label(
        log_card,
        text="Log da execução",
        font=(
            FONT,
            11,
            "bold",
        ),
        background=BG_CARD,
        foreground=TEXT_MAIN,
    ).grid(
        row=0,
        column=0,
        sticky="w",
    )

    log_box = tk.Text(
        log_card,
        wrap="word",
        font=(FONT, 9),
        bg="#FBFDFF",
        fg=TEXT_MAIN,
        relief="solid",
        bd=1,
        padx=8,
        pady=6,
    )

    log_box.grid(
        row=1,
        column=0,
        sticky="nsew",
        pady=(8, 10),
    )

    scroll = ttk.Scrollbar(
        log_card,
        command=log_box.yview,
    )

    scroll.grid(
        row=1,
        column=1,
        sticky="ns",
        pady=(8, 10),
    )

    log_box.configure(
        yscrollcommand=scroll.set,
        state="disabled",
    )

    # ==================================================================
    # FOOTER
    # ==================================================================

    footer = ttk.Frame(
        root
    )

    footer.grid(
        row=3,
        column=0,
        sticky="ew",
        padx=24,
        pady=(0, 14),
    )

    footer.grid_columnconfigure(
        0,
        weight=1,
    )

    ttk.Label(
        footer,
        textvariable=status_var,
        background=BG_MAIN,
        foreground=TEXT_MUTED,
        font=(FONT, 9),
    ).grid(
        row=0,
        column=0,
        sticky="w",
    )

    btn = ttk.Button(
        footer,
        text="Executar",
        style="Accent.TButton",
        width=14,
    )

    btn.grid(
        row=0,
        column=1,
        sticky="e",
        padx=(12, 0),
    )

    # ==================================================================
    # LOG
    # ==================================================================

    def log(msg):

        def _append():

            log_box.configure(
                state="normal"
            )

            log_box.insert(
                "end",
                f"{datetime.now().strftime('%H:%M:%S')}  "
                f"{msg}\n",
            )

            log_box.see(
                "end"
            )

            log_box.configure(
                state="disabled"
            )

        root.after(
            0,
            _append,
        )

    def atualizar_status(texto):

        root.after(
            0,
            lambda:
            status_var.set(
                texto
            ),
        )

    def finalizar_botao():

        root.after(
            0,
            lambda:
            btn.configure(
                state="normal"
            ),
        )

    # ==================================================================
    # VALIDAÇÃO
    # ==================================================================

    def validar_entrada():

        try:

            mes = int(
                mes_var.get()
            )

            ano = int(
                ano_var.get()
            )

            if not 1 <= mes <= 12:
                raise ValueError

            if not 1900 <= ano <= 2500:
                raise ValueError

        except Exception:

            raise ValueError(
                "Informe um mês e ano válidos.\n"
                "Exemplo: 09 / 2026"
            )

        if not pasta_var.get().strip():

            raise ValueError(
                "Informe a pasta de saída."
            )

        if not nome_var.get().strip():

            raise ValueError(
                "Informe o nome do arquivo de saída."
            )

        modo = modo_var.get().strip().upper()

        if modo not in ("CPF", "PIS"):
            raise ValueError(
                "Selecione CPF ou PIS em 'Agrupar por'."
            )

        return mes, ano, modo

    # ==================================================================
    # EXECUTAR
    # ==================================================================

    def executar():

        try:

            mes, ano, modo = (
                validar_entrada()
            )

        except ValueError as exc:

            messagebox.showwarning(
                "Validação",
                str(exc),
                parent=root,
            )

            return

        nome_arquivo = nome_var.get().strip()

        # garante extensão .xlsx
        if not nome_arquivo.lower().endswith(".xlsx"):
            nome_arquivo += ".xlsx"

        caminho_saida = os.path.join(
            pasta_var.get().strip(),
            nome_arquivo,
        )

        if not os.path.isabs(
            caminho_saida
        ):

            caminho_saida = os.path.abspath(
                caminho_saida
            )

        btn.configure(
            state="disabled"
        )

        log_box.configure(
            state="normal"
        )

        log_box.delete(
            "1.0",
            "end",
        )

        log_box.configure(
            state="disabled"
        )

        atualizar_status(
            "Executando..."
        )

        log(
            "Início do processamento."
        )

        log(
            f"Período solicitado: "
            f"{mes:02d}/{ano}."
        )

        log(
            f"Modo de agrupamento: {modo}."
        )

        if modo == "CPF":
            log(
                "Regra CPF: somente 11 dígitos são aceitos; "
                "demais formatos vão para Inválidos."
            )
        else:
            log(
                "Regra PIS: 11 dígitos mantidos; "
                "12 dígitos perdem o primeiro; "
                "demais formatos vão para Inválidos."
            )

        log(
            "Regra vínculo: "
            "consulta primeiro; "
            "prefixo da matrícula como fallback."
        )

        def worker():

            try:

                os.makedirs(
                    os.path.dirname(
                        caminho_saida
                    )
                    or ".",
                    exist_ok=True,
                )

                df = consultar_banco(
                    log
                )

                executar_calculo_df(
                    df,
                    mes,
                    ano,
                    caminho_saida,
                    log,
                    modo,
                )

                atualizar_status(
                    "Concluído com sucesso."
                )

                root.after(
                    0,
                    lambda:
                    messagebox.showinfo(
                        "Quinquênios",
                        "Processamento concluído.\n\n"
                        "Arquivo gerado em:\n"
                        f"{caminho_saida}",
                        parent=root,
                    ),
                )

            except Exception as exc:

                log(
                    f"ERRO: {exc}"
                )

                atualizar_status(
                    "Erro durante a execução."
                )

                root.after(
                    0,
                    lambda:
                    messagebox.showerror(
                        "Erro",
                        str(exc),
                        parent=root,
                    ),
                )

            finally:

                finalizar_botao()

        threading.Thread(
            target=worker,
            daemon=True,
        ).start()

    btn.configure(
        command=executar
    )

    centralizar(
        root
    )

    return root

    # ==================================================================
    # TESTES ESPECÍFICOS DA REGRA DE ANO/MÊS + DATA-BASE
    # ==================================================================

    linha += 2

    ws.merge_cells(
        start_row=linha,
        start_column=1,
        end_row=linha,
        end_column=12,
    )

    c = ws.cell(
        linha,
        1,
        "TESTES CRÍTICOS — DIA IGNORADO x MÊS/ANO RESPEITADO",
    )

    c.font = Font(
        size=12,
        bold=True,
        color=branco,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=vermelho,
    )

    c.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    linha += 1

    cab_criticos = [
        "Teste",
        "Início",
        "Fim",
        "Meses Python",
        "Anos Python",
        "Resultado esperado",
        "Status",
    ]

    for col, titulo in enumerate(
        cab_criticos,
        start=1,
    ):

        if col == 7:
            ws.merge_cells(
                start_row=linha,
                start_column=col,
                end_row=linha,
                end_column=12,
            )

        c = ws.cell(
            linha,
            col,
            titulo,
        )

        c.font = Font(
            bold=True,
            color=branco,
        )

        c.fill = PatternFill(
            "solid",
            fgColor=azul_escuro,
        )

        c.border = borda

        c.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    linha += 1

    testes_criticos = [
        {
            "nome":
                "Mesmo mês/ano — dia anterior",
            "inicio":
                date(2012, 6, 15),
            "fim":
                date(2026, 6, 14),
            "meses_esperados":
                168,
            "anos_esperados":
                14,
            "resultado_esperado":
                "14 anos — CORRETO",
        },
        {
            "nome":
                "Mês anterior ao aniversário",
            "inicio":
                date(2021, 11, 1),
            "fim":
                date(2026, 10, 30),
            "meses_esperados":
                59,
            "anos_esperados":
                4,
            "resultado_esperado":
                "4 anos e 11 meses — NÃO completa 5",
        },
        {
            "nome":
                "Mesmo mês do aniversário",
            "inicio":
                date(2021, 11, 1),
            "fim":
                date(2026, 11, 30),
            "meses_esperados":
                60,
            "anos_esperados":
                5,
            "resultado_esperado":
                "5 anos — COMPLETA",
        },
    ]

    for teste in testes_criticos:

        inicio_t = teste["inicio"]
        fim_t = teste["fim"]

        meses_python = meses_intervalo(
            inicio_t,
            fim_t,
        )

        anos_python = (
            meses_python // 12
        )

        esperado_ok = (
            meses_python
            == teste["meses_esperados"]
            and
            anos_python
            == teste["anos_esperados"]
        )

        valores = [
            teste["nome"],
            inicio_t,
            fim_t,
            meses_python,
            anos_python,
            teste["resultado_esperado"],
            "OK" if esperado_ok else "ERRO",
        ]

        for col, valor in enumerate(
            valores,
            start=1,
        ):

            if col == 7:
                ws.merge_cells(
                    start_row=linha,
                    start_column=7,
                    end_row=linha,
                    end_column=12,
                )

            c = ws.cell(
                linha,
                col,
                valor,
            )

            c.border = borda

            c.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            if col in (2, 3):
                c.number_format = "dd/mm/yyyy"

            if col == 7:
                c.fill = PatternFill(
                    "solid",
                    fgColor=(
                        verde_claro
                        if esperado_ok
                        else vermelho_claro
                    ),
                )

                c.font = Font(
                    bold=True,
                    color=(
                        verde
                        if esperado_ok
                        else vermelho
                    ),
                )

            else:
                c.fill = PatternFill(
                    "solid",
                    fgColor=(
                        azul_muito_claro
                        if linha % 2 == 0
                        else branco
                    ),
                )

        linha += 1
# ===========================================================================
# MAIN
# ===========================================================================

if __name__ == "__main__":

    app = montar_gui()

    app.mainloop()
