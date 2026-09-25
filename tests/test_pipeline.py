"""Teste de ponta a ponta: SQLite sintético -> consulta -> cálculo -> Excel."""

import openpyxl
import pytest

import generate_synthetic_data as synth
import main

MES, ANO = 10, 2026


@pytest.fixture(scope="module")
def df(tmp_path_factory):
    caminho = tmp_path_factory.mktemp("db") / "synthetic_hr.db"
    synth.build_database(str(caminho))
    main.DB_PATH = str(caminho)
    return main.consultar_banco()


def _executar(df, modo, tmp_path):
    saida = tmp_path / f"quinquenios_{modo}.xlsx"
    main.executar_calculo_df(df, MES, ANO, str(saida), None, modo)
    wb = openpyxl.load_workbook(saida)
    selecionados = {
        row[2]: row[8]
        for row in wb["Quinquenios"].iter_rows(min_row=2, values_only=True)
        if row[2]
    }
    invalidos = {
        row[0]
        for row in wb["Inválidos"].iter_rows(values_only=True)
        if row[0] and str(row[0]).startswith("Employee")
    }
    return wb, selecionados, invalidos


def test_consulta_aplica_filtros_da_query(df):
    assert len(df) == 51
    assert not set(df["MATRÍCULA"]) & {"99999901", "99999902"}
    assert not df["NOME_FUNCIONÁRIO"].str.contains("903").any()  # empresa CO3
    assert set(df["EMPRESA"]) <= {"CO1", "CO2"}


def test_consulta_devolve_colunas_esperadas(df):
    assert set(main.COLUNAS_ESPERADAS) <= set(df.columns)


@pytest.mark.parametrize("modo", ["CPF", "PIS"])
def test_resultado_por_modo(df, tmp_path, modo):
    wb, selecionados, invalidos = _executar(df, modo, tmp_path)
    esperado = synth.EXPECTED[modo]

    assert selecionados == esperado["selected"]
    assert invalidos == esperado["invalid_names"]
    # abas de auditoria preservadas
    assert {
        "Detalhe por Matricula",
        "Detalhe por Mat. Geral",
        "Inválidos",
        "Debug",
    } <= set(wb.sheetnames)


def test_pessoa_com_tres_matriculas_usa_a_ativa_como_principal(df, tmp_path):
    wb, _, _ = _executar(df, "CPF", tmp_path)
    linhas = [
        r
        for r in wb["Quinquenios"].iter_rows(min_row=2, values_only=True)
        if r[2] == "Employee 008"
    ]
    assert len(linhas) == 1
    matricula, vinculo = linhas[0][0], linhas[0][1]
    assert matricula.startswith("04")     # prefixo da matrícula ativa
    assert vinculo == "ENTITY C (CLT)"    # vínculo vindo da consulta


def test_gerador_e_deterministico(tmp_path):
    a = tmp_path / "a.db"
    b = tmp_path / "b.db"
    synth.build_database(str(a))
    synth.build_database(str(b))
    main.DB_PATH = str(a)
    df_a = main.consultar_banco()
    main.DB_PATH = str(b)
    df_b = main.consultar_banco()
    assert df_a.equals(df_b)


def test_sql_publico_em_arquivo_espelha_o_do_codigo():
    import os

    raiz = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    with open(os.path.join(raiz, "sql", "quintennial_query.sql"), encoding="utf-8") as f:
        conteudo = "".join(l for l in f if not l.startswith("--"))
    assert conteudo.strip() == main.SQL_QUINQUENIOS.strip()
